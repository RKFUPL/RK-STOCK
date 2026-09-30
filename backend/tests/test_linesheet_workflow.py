from io import BytesIO

from openpyxl import load_workbook

from app.routes import _build_linesheet_excel
from app.sku import inventory_sku, linesheet_sku
from app.utils import now


def test_centralized_linesheet_and_inventory_skus():
    base = linesheet_sku("Inaara", "RK104", "Red", 2)
    assert base == "RK-INA-RK104-RED-2"
    assert inventory_sku(base, "XS") == "RK-INA-RK104-RED-2-XS"
    assert inventory_sku(base, "XL") == "RK-INA-RK104-RED-2-XL"
    assert "-XS" not in base


def test_simplified_linesheet_draft_persists_without_stock(client, headers, app):
    response = client.post("/api/linesheets", headers=headers, json={"name": "Inaara Preview", "collection": "Inaara", "items": [{"product_code": "RK104", "description": "Embroidered kurta set", "color": "Red", "set_of": 2, "sizes": ["XS", "M", "XL"], "mrp": 12500}]})
    assert response.status_code == 201
    detail = client.get(f"/api/linesheets/{response.json['id']}", headers=headers).json["linesheet"]
    assert detail["status"] == "draft"
    assert detail.get("client_id") is None
    assert detail["items"][0]["linesheet_sku"] == "RK-INA-RK104-RED-2"
    assert detail["items"][0]["sizes"] == ["XS", "M", "XL"]
    assert "po_folder" not in detail
    assert client.get(f"/api/linesheets/{response.json['id']}/purchase-orders", headers=headers).status_code == 409
    assert app.extensions["mongo_db"].stock_balances.count_documents({}) == 0
    assert client.get(f"/api/linesheets/{response.json['id']}/export.xlsx", headers=headers).status_code == 200
    pdf = client.get(f"/api/linesheets/{response.json['id']}/export.pdf", headers=headers)
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF")


def test_duplicate_linesheet_sku_in_one_draft_is_rejected(client, headers):
    item = {"product_code": "RK104", "description": "Kurta", "color": "Red", "set_of": 2, "sizes": ["M"], "mrp": 1000}
    response = client.post("/api/linesheets", headers=headers, json={"name": "Duplicate test", "collection": "Inaara", "items": [item, item]})
    assert response.status_code == 400
    assert "Duplicate product configuration" in response.json["error"]


def test_inventory_variants_are_size_specific_and_reused(client, headers, app):
    sheet = client.post("/api/linesheets", headers=headers, json={"name": "Inventory source", "collection": "Inaara", "items": [{"product_code": "RK200", "description": "Jacket", "color": "Black", "set_of": 1, "sizes": ["S", "M"], "mrp": 5000}]}).json
    detail = client.get(f"/api/linesheets/{sheet['id']}", headers=headers).json["linesheet"]
    base = detail["items"][0]["linesheet_sku"]
    first = client.post("/api/inventory/variants", headers=headers, json={"linesheet_sku": base, "sizes": ["S", "M"], "quantities": {"S": 3, "M": 4}, "request_id": "first"})
    assert first.status_code == 201
    assert first.json["created"] == [f"{base}-S", f"{base}-M"]
    second = client.post("/api/inventory/variants", headers=headers, json={"linesheet_sku": base, "sizes": ["S", "M"], "quantities": {}})
    assert second.status_code == 201
    assert second.json["created"] == []
    assert set(second.json["reused"]) == {f"{base}-S", f"{base}-M"}
    balances = list(app.extensions["mongo_db"].stock_balances.find({}))
    assert {(row["sku"], row["physical"]) for row in balances} == {(f"{base}-S", 3), (f"{base}-M", 4)}


def test_same_named_mds_linesheets_get_distinct_id_backed_po_folders(client, headers, app):
    database = app.extensions["mongo_db"]
    ids = [database.linesheets.insert_one({"name": "Aakaar July", "linesheet_number": f"LS-{index}", "type": "mds_outright", "items": [], "status": "draft", "created_at": now()}).inserted_id for index in (1, 2)]
    folders = [client.get(f"/api/linesheets/{item}/purchase-orders", headers=headers).json["folder"] for item in ids]
    assert folders[0]["id"] != folders[1]["id"]
    assert folders[0]["display_name"] == folders[1]["display_name"] == "Aakaar July"


def test_linesheet_detail_and_export_normalize_aakaar_source_order_quantity(client, headers, app):
    database = app.extensions["mongo_db"]
    sheet_id = database.linesheets.insert_one({
        "name": "Aakaar compatibility fixture",
        "linesheet_number": "AAKAAR-COMPAT",
        "collection": "Aakaar",
        "items": [
            {"product_code": "CK-201", "vendor_code": "CK-201-Red", "color": "Red", "size_quantities": {"M": 3}, "sizes": ["M"], "mrp": "100.00"},
            {"product_code": "CK-202", "source_vendor_code": "CK-202-Ivory", "color": "Ivory", "source_order_quantity": {"M": 1}, "sizes": ["M"], "mrp": "200.00"},
        ],
        "status": "draft",
        "created_at": now(),
    }).inserted_id

    detail = client.get(f"/api/linesheets/{sheet_id}", headers=headers)
    assert detail.status_code == 200
    items = detail.json["linesheet"]["items"]
    assert items[0]["size_quantities"] == {"M": 3}
    assert items[1]["size_quantities"] == {"M": 1}
    assert items[1]["source_order_quantity"] == {"M": 1}
    assert items[0]["vendor_code"] == "CK-201-Red"
    assert items[1]["vendor_code"] == "CK-202-Ivory"

    workbook_bytes = _build_linesheet_excel(database.linesheets.find_one({"_id": sheet_id}))
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    vendor_code_column = headers.index("Vendor Code") + 1
    mrp_column = headers.index("MRP") + 1
    assert "M" in headers
    assert worksheet.cell(row=2, column=vendor_code_column).value == "CK-201-Red"
    assert worksheet.cell(row=3, column=vendor_code_column).value == "CK-202-Ivory"
    assert worksheet.cell(row=2, column=mrp_column).value == 100
    assert worksheet.cell(row=3, column=mrp_column).value == 200
    assert worksheet.cell(row=2, column=headers.index("M") + 1).value == 3
    assert worksheet.cell(row=3, column=headers.index("M") + 1).value == 1


def test_mds_workdrive_sync_uploads_generated_files_and_po(client, headers, app, monkeypatch):
    database = app.extensions["mongo_db"]
    sheet_id = database.linesheets.insert_one({"name": "MDS WorkDrive", "linesheet_number": "LS-WD-1", "type": "mds_outright", "items": [], "status": "draft", "created_at": now()}).inserted_id
    uploads = []
    class Stub:
        configured = True
        root_folder_id = "root"
        def ensure_managed_folder(self, key, name, parent_id): return {"mds-root": "mds", f"mds-linesheet:{sheet_id}": "sheet", f"mds-linesheet:{sheet_id}:purchase-orders": "po"}[key]
        def upload_bytes(self, content, parent_id, filename, content_type):
            uploads.append((parent_id, filename, content_type, content))
            return {"data": {"id": f"file-{len(uploads)}", "attributes": {"permalink": "https://example.invalid/file"}}}
    monkeypatch.setattr("app.routes.WorkDriveClient", Stub)
    synced = client.post(f"/api/linesheets/{sheet_id}/sync-workdrive", headers=headers)
    assert synced.status_code == 200
    assert [item[2] for item in uploads] == ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/pdf"]
    po = client.post(f"/api/linesheets/{sheet_id}/purchase-orders", headers=headers, data={"files": (__import__('io').BytesIO(b"%PDF-1.7 PO"), "po.pdf")}, content_type="multipart/form-data")
    assert po.status_code == 201
    assert po.json["items"][0]["workdrive_folder_id"] == "po"
    assert po.json["items"][0]["upload_status"] == "synced"
