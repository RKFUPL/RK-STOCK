from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from cryptography.fernet import Fernet
from openpyxl import Workbook, load_workbook
import pytest

from app.utils import utc_datetime
from app.workdrive import WorkDriveClient, WorkDriveError


def create_client(client, headers):
    response = client.post("/api/clients", headers=headers, json={"name": "MDS Delhi", "client_code": "MDS-DL", "category": "mds", "contact_person": "Test Contact", "email": "mds@example.com", "phone": "9999999999", "address": "Test address"})
    assert response.status_code == 201
    return response.json["id"]


def create_order(client, headers, client_id):
    response = client.post("/api/orders", headers=headers, json={"client_id": client_id, "order_type": "mds_outright", "items": [{"sku": "RK-001", "product_name": "Inaara", "color": "Red", "size": "M", "quantity": 10, "unit_price": 2500}]})
    assert response.status_code == 201
    detail = client.get(f"/api/orders/{response.json['id']}", headers=headers).json
    return response.json["id"], detail["order"]["items"][0]["line_id"]


def test_dashboard_is_database_driven_and_empty(client, headers):
    response = client.get("/api/dashboard?period=all", headers=headers)
    assert response.status_code == 200
    assert response.json["sections"]["mds_outright"]["received"] == 0
    client_id = create_client(client, headers)
    create_order(client, headers, client_id)
    response = client.get("/api/dashboard?period=all", headers=headers)
    assert response.json["sections"]["mds_outright"]["received"] == 1


def test_partial_production_movement_conserves_quantity(client, headers):
    client_id = create_client(client, headers)
    order_id, line_id = create_order(client, headers, client_id)
    response = client.post(f"/api/orders/{order_id}/production", headers=headers, json={"line_id": line_id, "quantity": 10})
    assert response.status_code == 201
    batch_id = response.json["id"]
    response = client.post(f"/api/production/{batch_id}/move", headers=headers, json={"from_stage": "designing", "to_stage": "stitching", "quantity": 4})
    assert response.status_code == 200
    buckets = response.json["batch"]["stage_quantities"]
    assert buckets["designing"] == 6
    assert buckets["stitching"] == 4
    assert sum(buckets.values()) == 10
    rejected = client.post(f"/api/production/{batch_id}/move", headers=headers, json={"from_stage": "stitching", "to_stage": "embroidery", "quantity": 5})
    assert rejected.status_code == 400


def test_idempotent_dispatch_does_not_double_deduct(client, headers, app):
    client_id = create_client(client, headers)
    order_id, line_id = create_order(client, headers, client_id)
    received = client.post("/api/stock", headers=headers, json={"sku": "RK-001", "color": "Red", "size": "M", "quantity": 10, "transaction_type": "opening_stock", "idempotency_key": "open-1"})
    assert received.status_code == 201
    assert client.post(f"/api/orders/{order_id}/reserve", headers=headers, json={"line_id": line_id, "quantity": 5}).status_code == 200
    dispatch_headers = {**headers, "Idempotency-Key": "dispatch-1"}
    first = client.post(f"/api/orders/{order_id}/dispatch", headers=dispatch_headers, json={"line_id": line_id, "quantity": 5})
    second = client.post(f"/api/orders/{order_id}/dispatch", headers=dispatch_headers, json={"line_id": line_id, "quantity": 5})
    assert first.status_code == 200 and first.json["created"] is True
    assert second.status_code == 400  # outstanding validation stops a replay before mutation
    database = app.extensions["mongo_db"]
    balance = database.stock_balances.find_one({"sku": "RK-001", "color": "Red", "size": "M", "location": "main"})
    assert balance["physical"] == 5
    assert balance["reserved"] == 0
    assert database.stock_ledger.count_documents({"transaction_id": "dispatch-1"}) == 1


def test_po_pdf_versions_are_retained(client, headers):
    client_id = create_client(client, headers)
    order_id, _ = create_order(client, headers, client_id)
    pdf = b"%PDF-1.4\n%%EOF"
    first = client.post(f"/api/orders/{order_id}/documents", headers=headers, data={"file": (BytesIO(pdf), "po.pdf")}, content_type="multipart/form-data")
    assert first.status_code == 201
    second = client.post(f"/api/orders/{order_id}/documents", headers=headers, data={"family_id": first.json["family_id"], "file": (BytesIO(pdf + b"\n"), "po-revised.pdf")}, content_type="multipart/form-data")
    assert second.status_code == 201 and second.json["version"] == 2
    listing = client.get(f"/api/orders/{order_id}/documents", headers=headers)
    assert [item["version"] for item in listing.json["items"]] == [2, 1]


def test_excel_linesheet_import_dynamic_sizes_and_export(client, headers):
    client_id = create_client(client, headers)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Outright Order Details"
    sheet.append(["Sr No", "Vendor Code", "Sku", "Color", "Category", "No of Components", "MRP", "L (38 inch)", "XL (40 inch)", "Remark", "Po DeliveryDate"])
    sheet.append([1, "0007", "RK-EXCEL-001", "Ivory", "Dress", 2, 12500, 3, 4, "Rush", "2026-10-20"])
    source = BytesIO(); workbook.save(source); source.seek(0)
    inspected = client.post("/api/linesheets/import/inspect", headers=headers, data={"file": (source, "outright.xlsx")}, content_type="multipart/form-data")
    assert inspected.status_code == 200
    assert [item["name"] for item in inspected.json["size_columns"]] == ["L", "XL"]
    assert inspected.json["rows"][0]["sku"] == "RK-EXCEL-001"
    committed = client.post(f"/api/linesheets/import/{inspected.json['import_id']}/commit", headers=headers, json={"client_id": client_id, "collection": "Inaara", "type": "mds_outright", "name": "Inaara October", "status": "draft"})
    assert committed.status_code == 201
    detail = client.get(f"/api/linesheets/{committed.json['id']}", headers=headers)
    assert detail.json["linesheet"]["po_folder"]["display_name"] == "Inaara October"
    item = detail.json["linesheet"]["items"][0]
    assert item["vendor_code"] == "0007"
    assert item["size_quantities"] == {"L": 3, "XL": 4}
    exported = client.get(f"/api/linesheets/{committed.json['id']}/export.xlsx", headers=headers)
    assert exported.status_code == 200
    output = load_workbook(BytesIO(exported.data), data_only=True)
    assert output.active.freeze_panes == "A2"
    assert "Sku" in [cell.value for cell in output.active[1]]
    po_upload = client.post(f"/api/linesheets/{committed.json['id']}/purchase-orders", headers=headers, data={"files": [(BytesIO(b"%PDF-1.7 first"), "first.pdf"), (BytesIO(b"%PDF-1.7 second"), "second.pdf")]}, content_type="multipart/form-data")
    assert po_upload.status_code == 201
    assert len(po_upload.json["items"]) == 2
    po_listing = client.get(f"/api/linesheets/{committed.json['id']}/purchase-orders", headers=headers)
    assert len(po_listing.json["items"]) == 2
    folder_before = po_listing.json["folder"]["id"]
    assert client.get(f"/api/linesheets/{committed.json['id']}/purchase-orders", headers=headers).json["folder"]["id"] == folder_before


def test_client_outright_fixture_parses_decorated_quantities(client, headers):
    fixtures = list(Path("backend/uploads/imports").glob("*/MTc4*.xlsx"))
    if not fixtures:
        pytest.skip("uploaded Outright fixture is not present")
    with fixtures[0].open("rb") as handle:
        response = client.post("/api/linesheets/import/inspect", headers=headers, data={"file": (handle, "outright-client.xlsx")}, content_type="multipart/form-data")
    assert response.status_code == 200
    payload = response.json
    assert payload["worksheet"] == "OutRight Order Details"
    assert payload["summary"]["total"] == 47
    assert payload["summary"]["invalid"] == 0
    assert all(row["size_quantities"].get("M") == 1 for row in payload["rows"])
    assert any("OR-" in reference for row in payload["rows"] for reference in row.get("references", []))


def test_workdrive_status_requires_auth_and_reports_configuration(client, headers):
    unauthenticated = client.get("/api/workdrive/status")
    assert unauthenticated.status_code == 401
    authenticated = client.get("/api/workdrive/status", headers=headers)
    assert authenticated.status_code == 200
    assert authenticated.json["verified"] is False
    assert "configured" in authenticated.json


def test_workdrive_managed_folder_is_idempotent(app, monkeypatch):
    calls = []
    monkeypatch.setattr(WorkDriveClient, "create_folder", lambda self, name, parent_id: calls.append((name, parent_id)) or {"id": "folder-1"})
    with app.app_context():
        workdrive = WorkDriveClient()
        assert workdrive.ensure_managed_folder("test-folder", "Test", "root") == "folder-1"
        assert workdrive.ensure_managed_folder("test-folder", "Test", "root") == "folder-1"
    assert calls == [("Test", "root")]


class _WorkDriveCallbackStub:
    def exchange_code(self, code):
        assert code == "oauth-code"

    def discover_root_folder(self):
        return {"id": "verified-folder", "name": "STOCK&LINESHEET", "team_id": "team-id"}


@pytest.mark.parametrize(
    "expires_at",
    [
        datetime.now(timezone.utc) + timedelta(minutes=5),
        datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(minutes=5),
    ],
    ids=["timezone-aware", "mongo-naive"],
)
def test_workdrive_callback_accepts_unexpired_naive_and_aware_state(client, app, monkeypatch, expires_at):
    monkeypatch.setattr("app.routes.WorkDriveClient", _WorkDriveCallbackStub)
    state = "valid-state"
    app.extensions["mongo_db"].settings.insert_one(
        {"_id": f"workdrive_state:{state}", "state": state, "user_id": "admin", "expires_at": expires_at}
    )

    response = client.get(f"/api/integrations/zoho/callback?state={state}&code=oauth-code")

    assert response.status_code == 302
    assert "workdrive=connected" in response.location
    assert app.extensions["mongo_db"].settings.find_one({"_id": f"workdrive_state:{state}"}) is None
    oauth = app.extensions["mongo_db"].settings.find_one({"_id": "workdrive_oauth"})
    assert oauth["root_folder_id"] == "verified-folder"
    assert oauth["verified_folder_id"] == "verified-folder"


def test_workdrive_callback_rejects_expired_state(client, app, monkeypatch):
    monkeypatch.setattr("app.routes.WorkDriveClient", _WorkDriveCallbackStub)
    state = "expired-state"
    app.extensions["mongo_db"].settings.insert_one(
        {
            "_id": f"workdrive_state:{state}",
            "state": state,
            "user_id": "admin",
            "expires_at": datetime.now(timezone.utc) - timedelta(minutes=1),
        }
    )

    response = client.get(f"/api/integrations/zoho/callback?state={state}&code=oauth-code")

    assert response.status_code == 302
    assert "Invalid%20or%20expired%20OAuth%20state" in response.location
    assert app.extensions["mongo_db"].settings.find_one({"_id": f"workdrive_state:{state}"}) is not None


def test_utc_datetime_normalizes_naive_and_aware_values():
    naive = datetime(2026, 1, 1, 12, 0, 0)
    aware = datetime(2026, 1, 1, 12, 0, 0, tzinfo=timezone(timedelta(hours=5, minutes=30)))

    assert utc_datetime(naive).tzinfo == timezone.utc
    assert utc_datetime(naive).hour == 12
    assert utc_datetime(aware) == datetime(2026, 1, 1, 6, 30, tzinfo=timezone.utc)


class _TokenResponse:
    ok = True
    status_code = 200

    def __init__(self, payload):
        self.payload = payload

    def json(self):
        return self.payload


def _configure_workdrive_token_test(monkeypatch, app):
    key = Fernet.generate_key().decode()
    monkeypatch.setenv("ZOHO_TOKEN_ENCRYPTION_KEY", key)
    monkeypatch.setenv("ZOHO_CLIENT_ID", "client-id")
    monkeypatch.setenv("ZOHO_CLIENT_SECRET", "client-secret")
    monkeypatch.setenv("ZOHO_REDIRECT_URI", "http://localhost:5006/api/integrations/zoho/callback")
    monkeypatch.setenv("ZOHO_WORKDRIVE_TEAM_ID", "team-id")
    monkeypatch.setenv("ZOHO_WORKDRIVE_ROOT_FOLDER_ID", "team-folder-123")
    with app.app_context():
        return key


def test_exchange_code_stores_encrypted_refresh_token(app, monkeypatch):
    key = _configure_workdrive_token_test(monkeypatch, app)
    monkeypatch.setattr("app.workdrive.requests.post", lambda *args, **kwargs: _TokenResponse({"access_token": "access", "refresh_token": "new-refresh"}))

    with app.app_context():
        client = WorkDriveClient()
        payload = client.exchange_code("oauth-code")
        stored = app.extensions["mongo_db"].settings.find_one({"_id": "workdrive_oauth"})
        assert payload["access_token"] == "access"
        assert stored["refresh_token_encrypted"] != "new-refresh"
        assert Fernet(key.encode()).decrypt(stored["refresh_token_encrypted"].encode()).decode() == "new-refresh"


def test_exchange_code_missing_refresh_token_preserves_existing_token(app, monkeypatch):
    key = _configure_workdrive_token_test(monkeypatch, app)
    monkeypatch.setattr("app.workdrive.requests.post", lambda *args, **kwargs: _TokenResponse({"access_token": "access"}))

    with app.app_context():
        client = WorkDriveClient()
        client.store_refresh_token("existing-refresh")
        before = app.extensions["mongo_db"].settings.find_one({"_id": "workdrive_oauth"})["refresh_token_encrypted"]
        payload = client.exchange_code("oauth-code")
        after = app.extensions["mongo_db"].settings.find_one({"_id": "workdrive_oauth"})["refresh_token_encrypted"]
        assert payload["access_token"] == "access"
        assert after == before
        assert Fernet(key.encode()).decrypt(after.encode()).decode() == "existing-refresh"


def test_exchange_code_missing_refresh_token_without_existing_token(app, monkeypatch):
    _configure_workdrive_token_test(monkeypatch, app)
    monkeypatch.setattr("app.workdrive.requests.post", lambda *args, **kwargs: _TokenResponse({"access_token": "access"}))

    with app.app_context():
        with pytest.raises(WorkDriveError, match="grant offline access") as error:
            WorkDriveClient().exchange_code("oauth-code")
        assert error.value.kind == "missing_refresh_token"


def test_authorization_and_token_exchange_use_canonical_redirect_uri(app, monkeypatch):
    _configure_workdrive_token_test(monkeypatch, app)
    calls = []

    def post(_url, **kwargs):
        calls.append(kwargs["data"])
        return _TokenResponse({"access_token": "access", "refresh_token": "refresh"})

    monkeypatch.setattr("app.workdrive.requests.post", post)
    with app.app_context():
        client = WorkDriveClient()
        authorization = parse_qs(urlparse(client.authorization_url("state")).query)
        client.exchange_code("oauth-code")

    assert authorization["redirect_uri"] == ["http://localhost:5006/api/integrations/zoho/callback"]
    assert calls[0]["redirect_uri"] == authorization["redirect_uri"][0]


def test_workdrive_environment_uses_canonical_names_and_rejects_placeholder_root(app, monkeypatch):
    _configure_workdrive_token_test(monkeypatch, app)
    monkeypatch.setenv("ZOHO_WORKDRIVE_ROOT_FOLDER_ID", "<folder-id>")

    with app.app_context():
        client = WorkDriveClient()
        assert client.client_id == "client-id"
        assert client.root_folder_configured is False
        assert "root_folder_id" in client.missing_configuration


def test_discovers_and_verifies_existing_team_folder(app, monkeypatch):
    _configure_workdrive_token_test(monkeypatch, app)
    monkeypatch.delenv("ZOHO_WORKDRIVE_ROOT_FOLDER_ID", raising=False)
    folder = {
        "id": "actual-folder-id",
        "type": "teamfolders",
        "attributes": {
            "display_attr_name": "STOCK&amp;LINESHEET",
            "parent_id": "team-id",
            "capabilities": {"can_read": True},
        },
    }

    with app.app_context():
        client = WorkDriveClient()
        client.store_refresh_token("refresh")

        monkeypatch.setattr("app.workdrive.requests.post", lambda *args, **kwargs: _TokenResponse({"access_token": "access"}))

        def get(url, **_kwargs):
            if url.endswith("/teams/team-id/teamfolders"):
                return _TokenResponse({"data": [folder]})
            if url.endswith("/teamfolders/actual-folder-id"):
                return _TokenResponse({"data": folder})
            raise AssertionError(f"Unexpected WorkDrive URL: {url}")

        monkeypatch.setattr("app.workdrive.requests.get", get)
        discovered = client.discover_root_folder()

    assert discovered == {"id": "actual-folder-id", "name": "STOCK&LINESHEET", "team_id": "team-id"}


def test_workdrive_status_reads_verified_folder_from_oauth_record(client, app, headers, monkeypatch):
    _configure_workdrive_token_test(monkeypatch, app)
    monkeypatch.delenv("ZOHO_WORKDRIVE_ROOT_FOLDER_ID", raising=False)
    with app.app_context():
        WorkDriveClient().store_refresh_token("refresh")
        app.extensions["mongo_db"].settings.update_one(
            {"_id": "workdrive_oauth"},
            {"$set": {"root_folder_id": "actual-folder-id", "verified_folder_id": "actual-folder-id", "verified_at": datetime.now(timezone.utc)}},
        )

    response = client.get("/api/workdrive/status", headers=headers)

    assert response.status_code == 200
    assert response.json["configured"] is True
    assert response.json["root_folder_configured"] is True
    assert response.json["verified"] is True
    assert response.json["status"] == "connected"
