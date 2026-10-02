from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import hashlib
import os
import re
from urllib.parse import quote

from bson import ObjectId
from flask import Blueprint, current_app, g, jsonify, request, send_file, redirect
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pymongo import DESCENDING, ReturnDocument
from pymongo.errors import DuplicateKeyError
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from werkzeug.utils import secure_filename

from .auth import auth_required, check_password, effective_permissions, hash_password, permission_required, token_for
from .db import db
from .services import INVENTORY_LOCATIONS, NON_SELLABLE_LOCATIONS, STAGES, activity, move_production, mutate_stock
from .fx import FxProvider, FxProviderError, SUPPORTED_CURRENCIES
from .sku import ALLOWED_SIZES, collection_code, inventory_sku, linesheet_sku
from .utils import money, next_number, now, oid, page_args, serialize, utc_datetime
from .workdrive import WorkDriveClient, WorkDriveError
from .mail import ZohoMailClient
from .storefront_integration import StorefrontIntegrationClient, StorefrontIntegrationError
from .aakaar_import import build_aakaar_projection, commit_aakaar_projection, existing_aakaar_identities
from .catalog_sync import CatalogSyncError, CatalogSyncService
from .cloudinary_service import CloudinaryConfigurationError, CloudinaryUploadError, delete_image, product_asset_folder, upload_image

api = Blueprint("api", __name__)


def body(required=()):
    data = request.get_json(silent=True) or {}
    missing = [field for field in required if data.get(field) in (None, "")]
    if missing:
        raise ValueError("Missing fields: " + ", ".join(missing))
    return data


def json_error(exc, code=400):
    return jsonify(error=str(exc)), code


def list_response(collection, query, sort=("created_at", DESCENDING)):
    page, limit = page_args(request)
    total = collection.count_documents(query)
    items = list(collection.find(query).sort(*sort).skip((page - 1) * limit).limit(limit))
    return jsonify(items=serialize(items), page=page, limit=limit, total=total)


def search_query(fields):
    query = {}
    value = request.args.get("q", "").strip()
    if value:
        query["$or"] = [{field: {"$regex": value, "$options": "i"}} for field in fields]
    for field in request.args:
        if field in fields and request.args[field]:
            query[field] = request.args[field]
    return query


def _is_mds_linesheet(sheet):
    return str(sheet.get("type") or "").startswith("mds_") or sheet.get("client_category") == "mds"


def _ensure_mds_po_folder(sheet, collection=None):
    if not _is_mds_linesheet(sheet):
        return None
    collection = collection or db().linesheets
    existing = sheet.get("po_folder") or {}
    folder_id = existing.get("id") or str(sheet["_id"])
    display_name = sheet.get("name") or sheet.get("title") or sheet.get("linesheet_number") or folder_id
    safe_name = secure_filename(display_name) or "linesheet"
    folder = (Path(current_app.config["UPLOAD_DIR"]) / "purchase-orders" / f"{folder_id}-{safe_name}").resolve()
    try:
        folder.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise RuntimeError("Unable to create the MDS purchase-order folder; retry the operation") from exc
    po_folder = {"id": folder_id, "linesheet_id": sheet["_id"], "display_name": display_name, "path": str(folder), "created_at": existing.get("created_at") or now()}
    collection.update_one({"_id": sheet["_id"]}, {"$set": {"po_folder": po_folder}})
    return po_folder


@api.get("/health")
def health():
    try:
        db().command("ping")
        return jsonify(status="ok", database="connected")
    except Exception:
        return jsonify(status="degraded", database="unavailable"), 503


@api.post("/auth/login")
def login():
    try:
        data = body(("password",))
        identifier = (data.get("identifier") or data.get("email") or "").strip()
        user = db().users.find_one({"$or": [{"email": identifier.lower()}, {"name": {"$regex": f"^{identifier}$", "$options": "i"}}], "active": True}) if identifier else None
        if not user or not check_password(data["password"], user["password_hash"]):
            return jsonify(error="Invalid email or password"), 401
        db().users.update_one({"_id": user["_id"]}, {"$set": {"last_login_at": now()}})
        return jsonify(token=token_for(user), user=serialize({k: v for k, v in user.items() if k != "password_hash"}))
    except ValueError as exc:
        return json_error(exc)


@api.get("/auth/me")
@auth_required
def me():
    payload = {k: v for k, v in g.user.items() if k != "password_hash"}
    payload["effective_permissions"] = sorted(effective_permissions(g.user))
    return jsonify(serialize(payload))


@api.route("/users", methods=["GET", "POST"])
@permission_required("users:write")
def users():
    if request.method == "GET":
        return list_response(db().users, {}, ("name", 1))
    try:
        data = body(("name", "email", "password", "role"))
        if data["role"] not in {"admin", "sales", "production_inventory"}:
            raise ValueError("Invalid role")
        doc = {"name": data["name"], "email": data["email"].strip().lower(), "password_hash": hash_password(data["password"]), "role": data["role"], "active": True, "created_at": now()}
        result = db().users.insert_one(doc)
        activity(g.user, "create", "user", result.inserted_id)
        return jsonify(id=str(result.inserted_id)), 201
    except Exception as exc:
        return json_error(exc)


@api.route("/clients", methods=["GET", "POST"])
@auth_required
def clients():
    if request.method == "GET":
        query = search_query(["name", "client_code", "category", "status"])
        return list_response(db().clients, query, ("name", 1))
    if g.user["role"] not in {"admin", "sales"}:
        return jsonify(error="Permission denied"), 403
    try:
        data = body(("name", "category", "contact_person", "email", "phone", "address"))
        if data["category"] not in {"mds", "direct"}:
            raise ValueError("Client category must be mds or direct")
        if "@" not in str(data["email"]):
            raise ValueError("A valid email address is required")
        client_code = str(data.get("client_code") or "").strip().upper() or None
        duplicate_query = [{"name": {"$regex": f"^{re.escape(data['name'].strip())}$", "$options": "i"}}]
        if client_code:
            duplicate_query.append({"client_code": client_code})
        if db().clients.find_one({"$or": duplicate_query}):
            raise ValueError("A client with this code or name already exists")
        doc = {key: data[key] for key in ("name", "category", "contact_person", "email", "phone", "address")}; doc.update({"client_code": client_code, "status": "active", "created_at": now(), "updated_at": now()})
        result = db().clients.insert_one(doc)
        activity(g.user, "create", "client", result.inserted_id)
        return jsonify(id=str(result.inserted_id)), 201
    except (ValueError, DuplicateKeyError) as exc:
        return json_error(exc)


@api.route("/clients/<client_id>", methods=["GET", "PATCH"])
@auth_required
def client_detail(client_id):
    client = db().clients.find_one({"_id": oid(client_id)})
    if not client:
        return jsonify(error="Client not found"), 404
    if request.method == "PATCH":
        if g.user["role"] not in {"admin", "sales"}:
            return jsonify(error="Permission denied"), 403
        changes = body()
        for protected in ("_id", "created_at"):
            changes.pop(protected, None)
        for required in ("name", "contact_person", "email", "phone", "address"):
            if required in changes and not str(changes[required]).strip():
                raise ValueError(f"{required} is required")
        if "email" in changes and "@" not in str(changes["email"]):
            raise ValueError("A valid email address is required")
        if "client_code" in changes:
            changes["client_code"] = str(changes["client_code"] or "").strip().upper() or None
        changes.pop("business_types", None)
        changes["updated_at"] = now()
        db().clients.update_one({"_id": client["_id"]}, {"$set": changes})
        activity(g.user, "update", "client", client_id, changes)
        client.update(changes)
    orders = list(db().orders.find({"client_id": client["_id"]}).sort("created_at", DESCENDING))
    documents = list(db().documents.find({"client_id": client["_id"]}).sort("created_at", DESCENDING))
    return jsonify(client=serialize(client), orders=serialize(orders), documents=serialize(documents))


@api.delete("/clients/<client_id>")
@auth_required
def delete_client(client_id):
    if g.user["role"] != "superadmin":
        return jsonify(error="Only superadmins can delete clients"), 403
    client = db().clients.find_one({"_id": oid(client_id)})
    if not client:
        return jsonify(error="Client not found"), 404
    if client.get("status") != "inactive":
        return jsonify(error="Client must be inactive before deletion"), 409
    linked = any(collection.count_documents({"client_id": client["_id"]}) for collection in (db().linesheets, db().orders, db().documents, db().client_linesheets))
    if linked:
        return jsonify(error="Client has associated linesheets, orders or documents; deactivate it instead"), 409
    db().clients.delete_one({"_id": client["_id"]}); activity(g.user, "delete", "client", client_id); return jsonify(ok=True)


@api.route("/clients/<client_id>/linesheets", methods=["GET", "POST"])
@auth_required
def client_linesheets(client_id):
    """Persistent client-uploaded linesheets, kept separate from RKFUPL sheets."""
    client = db().clients.find_one({"_id": oid(client_id)})
    if not client:
        return jsonify(error="Client not found"), 404
    if request.method == "GET":
        query = {"client_id": client["_id"], "kind": "client_linesheet"}
        value = request.args.get("q", "").strip()
        if value:
            query["$or"] = [{"title": {"$regex": value, "$options": "i"}}, {"original_filename": {"$regex": value, "$options": "i"}}]
        items = list(db().client_linesheets.find(query).sort("uploaded_at", DESCENDING))
        for item in items:
            item["purchase_order_count"] = db().documents.count_documents({"client_linesheet_id": item["_id"], "kind": "client_purchase_order"})
        return jsonify(items=serialize(items))
    if g.user["role"] not in {"admin", "sales"}:
        return jsonify(error="Permission denied"), 403
    if client.get("status") != "active":
        return jsonify(error="Inactive clients cannot receive new linesheets"), 409
    uploaded = request.files.get("file")
    title = (request.form.get("title") or "").strip()
    if not uploaded or not title:
        return jsonify(error="A title and linesheet file are required"), 400
    if uploaded.mimetype not in {"application/pdf", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel"}:
        return jsonify(error="Linesheet must be a PDF or Excel workbook"), 400
    data = uploaded.read()
    if not data:
        return jsonify(error="Uploaded linesheet is empty"), 400
    sheet_id = ObjectId()
    folder = (Path(current_app.config["UPLOAD_DIR"]) / "client-linesheets" / str(sheet_id)).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    filename = secure_filename(uploaded.filename or "linesheet") or "linesheet"
    path = folder / filename
    path.write_bytes(data)
    doc = {"_id": sheet_id, "kind": "client_linesheet", "client_id": client["_id"], "client_name": client["name"], "client_category": client.get("category"), "title": title, "original_filename": filename, "stored_path": str(path), "content_type": uploaded.mimetype, "size": len(data), "linesheet_type": request.form.get("type") or None, "uploaded_by": g.user["_id"], "uploaded_by_email": g.user["email"], "uploaded_at": now(), "created_at": now(), "updated_at": now()}
    db().client_linesheets.insert_one(doc)
    try:
        _ensure_mds_po_folder(doc, db().client_linesheets)
    except RuntimeError as exc:
        db().client_linesheets.delete_one({"_id": sheet_id})
        return jsonify(error=str(exc), retryable=True), 503
    if _is_mds_linesheet(doc):
        _sync_uploaded_mds_linesheet_workdrive(doc, data)
    activity(g.user, "upload_client_linesheet", "client_linesheet", sheet_id, {"client_id": client_id})
    return jsonify(id=str(sheet_id), item=serialize({key: value for key, value in doc.items() if key != "stored_path"})), 201


@api.get("/client-linesheets/<linesheet_id>/download")
@auth_required
def download_client_linesheet(linesheet_id):
    sheet = db().client_linesheets.find_one({"_id": oid(linesheet_id), "kind": "client_linesheet"})
    if not sheet or not os.path.exists(sheet.get("stored_path", "")):
        return jsonify(error="Client linesheet file not found"), 404
    return send_file(sheet["stored_path"], mimetype=sheet["content_type"], as_attachment=True, download_name=sheet["original_filename"])


@api.route("/client-linesheets/<linesheet_id>/purchase-orders", methods=["GET", "POST"])
@auth_required
def client_linesheet_purchase_orders(linesheet_id):
    sheet = db().client_linesheets.find_one({"_id": oid(linesheet_id), "kind": "client_linesheet"})
    if not sheet:
        return jsonify(error="Client linesheet not found"), 404
    po_folder = _ensure_mds_po_folder(sheet, db().client_linesheets) if _is_mds_linesheet(sheet) else None
    if request.method == "GET":
        return jsonify(items=serialize(list(db().documents.find({"client_linesheet_id": sheet["_id"], "kind": "client_purchase_order"}).sort("created_at", DESCENDING))))
    if g.user["role"] not in {"admin", "sales"}:
        return jsonify(error="Permission denied"), 403
    uploaded = request.files.get("file")
    if not uploaded or uploaded.mimetype != "application/pdf":
        return jsonify(error="A PDF purchase order is required"), 400
    data = uploaded.read()
    if not data.startswith(b"%PDF"):
        return jsonify(error="Uploaded file is not a valid PDF"), 400
    folder = Path(po_folder["path"]) if po_folder else (Path(current_app.config["UPLOAD_DIR"]) / "client-linesheets" / str(sheet["_id"]) / "purchase-orders").resolve()
    folder.mkdir(parents=True, exist_ok=True)
    document_id = ObjectId(); filename = secure_filename(uploaded.filename or "purchase-order.pdf") or "purchase-order.pdf"
    path = folder / f"{document_id}-{filename}"; path.write_bytes(data)
    doc = {"_id": document_id, "family_id": str(document_id), "version": 1, "kind": "client_purchase_order", "client_linesheet_id": sheet["_id"], "client_id": sheet["client_id"], "original_name": filename, "stored_path": str(path), "content_type": "application/pdf", "size": len(data), "sha256": hashlib.sha256(data).hexdigest(), "uploaded_by": g.user["_id"], "uploaded_by_email": g.user["email"], "created_at": now()}
    if _is_mds_linesheet(sheet):
        refreshed = db().client_linesheets.find_one({"_id": sheet["_id"]}) or sheet
        workdrive = WorkDriveClient()
        if workdrive.configured and refreshed.get("workdrive_po_folder_id"):
            try:
                item = _workdrive_uploaded_item(workdrive.upload_bytes(data, refreshed["workdrive_po_folder_id"], filename, "application/pdf"))
                if item.get("id"):
                    doc.update({"upload_status": "synced", "storage_provider": "local_and_workdrive", "workdrive_folder_id": refreshed["workdrive_po_folder_id"], "workdrive_file_id": item["id"], "workdrive_uploaded_at": now()})
            except WorkDriveError as exc:
                doc.update({"upload_status": "failed", "sync_error": str(exc), "sync_error_kind": exc.kind})
    db().documents.insert_one(doc)
    activity(g.user, "upload_client_purchase_order", "client_linesheet", linesheet_id, {"document_id": str(document_id)})
    return jsonify(id=str(document_id), item=serialize(doc)), 201


@api.route("/products", methods=["GET", "POST"])
@auth_required
def products():
    if request.method == "GET":
        return list_response(db().products, search_query(["name", "sku", "vendor_code", "category", "collection", "season"]), ("name", 1))
    if g.user["role"] != "admin":
        return jsonify(error="Permission denied"), 403
    try:
        data = body(("name", "sku"))
        sku = data["sku"].strip().upper()
        product_code = str(data.get("product_code") or sku).strip()
        collection_id = oid(data.get("collection_id")) if data.get("collection_id") else None
        if collection_id and not db().collections.find_one({"_id": collection_id, "active": True}):
            raise ValueError("A valid active collection is required")
        if collection_id and db().products.find_one({"collection_id": collection_id, "product_code": product_code}):
            raise ValueError("Product code already exists in this collection")
        doc = {**data, "sku": sku, "product_code": product_code, "collection_id": collection_id, "base_currency": str(data.get("base_currency") or "INR").upper(), "tax_inclusive": bool(data.get("tax_inclusive", True)), "cost_price": money(data.get("cost_price")), "selling_price": money(data.get("selling_price")), "colors": data.get("colors", []), "sizes": data.get("sizes", []), "active": data.get("active", True), "created_at": now(), "updated_at": now()}
        if doc["base_currency"] not in SUPPORTED_CURRENCIES:
            raise ValueError("Unsupported currency")
        result = db().products.insert_one(doc)
        activity(g.user, "create", "product", result.inserted_id)
        return jsonify(id=str(result.inserted_id)), 201
    except Exception as exc:
        return json_error(exc)


@api.patch("/products/<product_id>")
@permission_required("settings:write")
def update_product(product_id):
    product_oid = oid(product_id)
    if not product_oid:
        return jsonify(error="Invalid product ID"), 400
    product = db().products.find_one({"_id": product_oid})
    if not product:
        return jsonify(error="Product not found"), 404
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify(error="A JSON object is required"), 400
    allowed = {"name", "sku", "product_code", "color", "colors", "category", "description", "price", "selling_price", "base_currency", "currency", "tax_inclusive", "active", "status", "sizes", "collection_ids"}
    unknown = set(data) - allowed
    if unknown:
        return jsonify(error=f"Unsupported product fields: {', '.join(sorted(unknown))}"), 400

    changes = {}
    if "name" in data:
        if not isinstance(data["name"], str) or not data["name"].strip():
            return jsonify(error="Product name is required"), 400
        changes["name"] = data["name"].strip()
    if "sku" in data:
        if not isinstance(data["sku"], str) or not data["sku"].strip():
            return jsonify(error="SKU is required"), 400
        sku = data["sku"].strip().upper()
        if db().products.find_one({"sku": sku, "_id": {"$ne": product_oid}}):
            return jsonify(error="SKU already exists"), 409
        changes["sku"] = sku
    if "product_code" in data:
        if not isinstance(data["product_code"], str) or not data["product_code"].strip():
            return jsonify(error="Product code is required"), 400
        product_code = data["product_code"].strip()
        if db().products.find_one({"product_code": product_code, "_id": {"$ne": product_oid}}):
            return jsonify(error="Product code already exists"), 409
        changes["product_code"] = product_code
    if "color" in data or "colors" in data:
        colors = data.get("colors", [data.get("color")] if data.get("color") else [])
        if not isinstance(colors, list) or any(not isinstance(value, str) or not value.strip() for value in colors):
            return jsonify(error="Colours must be an array of non-empty strings"), 400
        changes["colors"] = list(dict.fromkeys(value.strip() for value in colors))
        if len(changes["colors"]) == 1:
            changes["color"] = changes["colors"][0]
    if "category" in data:
        category = data["category"]
        state = _category_state()
        if not isinstance(category, str) or category.strip().casefold() not in {value.casefold() for value in state["categories"]}:
            return jsonify(error="A valid configured category is required"), 400
        changes["category"] = next(value for value in state["categories"] if value.casefold() == category.strip().casefold())
    if "description" in data:
        if not isinstance(data["description"], str):
            return jsonify(error="Description must be a string"), 400
        changes["description"] = data["description"]
    for field in ("price", "selling_price"):
        if field in data:
            try:
                value = money(data[field])
            except Exception:
                return jsonify(error=f"{field} must be numeric"), 400
            if value is None or float(value) < 0:
                return jsonify(error=f"{field} must be non-negative"), 400
            changes[field] = value
    if "base_currency" in data or "currency" in data:
        currency = str(data.get("base_currency", data.get("currency", "INR"))).upper()
        if currency not in SUPPORTED_CURRENCIES:
            return jsonify(error="Unsupported currency"), 400
        changes["base_currency"] = currency
    if "tax_inclusive" in data:
        if not isinstance(data["tax_inclusive"], bool):
            return jsonify(error="Tax-inclusive state must be boolean"), 400
        changes["tax_inclusive"] = data["tax_inclusive"]
    if "active" in data:
        if not isinstance(data["active"], bool):
            return jsonify(error="Active state must be boolean"), 400
        changes["active"] = data["active"]
        changes["status"] = "active" if data["active"] else "archived"
    elif "status" in data:
        if data["status"] not in {"active", "inactive", "archived", "draft"}:
            return jsonify(error="Invalid product status"), 400
        changes["status"] = data["status"]
        changes["active"] = data["status"] == "active"
    if "sizes" in data:
        if not isinstance(data["sizes"], list) or any(not isinstance(value, str) or not value.strip() for value in data["sizes"]):
            return jsonify(error="Sizes must be an array of non-empty strings"), 400
        changes["sizes"] = list(dict.fromkeys(value.strip().upper() for value in data["sizes"]))
    if "collection_ids" in data:
        values = data["collection_ids"]
        if not isinstance(values, list):
            return jsonify(error="collection_ids must be an array"), 400
        collection_oids = []
        for value in values:
            item = oid(value)
            if not item or not db().collections.find_one({"_id": item, "active": True}):
                return jsonify(error="All collection_ids must reference active collections"), 400
            if item not in collection_oids:
                collection_oids.append(item)
        memberships = list(db().collections.find({"_id": {"$in": collection_oids}}))
        by_id = {item["_id"]: item for item in memberships}
        changes["collection_ids"] = collection_oids
        changes["collection_names"] = [by_id[item]["name"] for item in collection_oids]
        changes["collections"] = [{"id": item, "name": by_id[item]["name"], "slug": by_id[item].get("slug"), "code": by_id[item].get("code")} for item in collection_oids]
        changes["collection"] = changes["collection_names"][0] if len(collection_oids) == 1 else None
        changes["collection_id"] = collection_oids[0] if len(collection_oids) == 1 else None
    if not changes:
        return jsonify(error="No editable product fields supplied"), 400
    changes["updated_at"] = now()
    db().products.update_one({"_id": product_oid}, {"$set": changes})
    activity(g.user, "update", "product", product_oid, {"fields": sorted(changes.keys())})
    return jsonify(serialize(db().products.find_one({"_id": product_oid})))


def _product_dependencies(product):
    database = db()
    product_id, sku = product["_id"], product.get("sku")
    variants = list(database.inventory_variants.find({"product_id": product_id}, {"inventory_sku": 1}))
    skus = [value for value in [sku, *(item.get("inventory_sku") for item in variants)] if value]
    return {
        "variants": len(variants),
        "configurations": database.product_configurations.count_documents({"product_id": product_id}),
        "stock_balances": database.stock_balances.count_documents({"sku": {"$in": skus}}) if skus else 0,
        "stock_ledger": database.stock_ledger.count_documents({"sku": {"$in": skus}}) if skus else 0,
        "orders": database.orders.count_documents({"$or": [{"items.sku": {"$in": skus}}, {"items.product_id": product_id}]}),
        "linesheets": database.linesheets.count_documents({"$or": [{"items.product_id": product_id}, {"items.product_code": product.get("product_code")}]}),
        "imports": database.imports.count_documents({"$or": [{"product_id": product_id}, {"rows.sku": {"$in": skus}}, {"preview.sku": {"$in": skus}}]}) if skus else database.imports.count_documents({"product_id": product_id}),
        "documents": database.documents.count_documents({"$or": [{"product_id": product_id}, {"sku": {"$in": skus}}]}) if skus else database.documents.count_documents({"product_id": product_id}),
    }


@api.delete("/products/<product_id>")
@permission_required("settings:write")
def delete_product(product_id):
    product_oid = oid(product_id)
    if not product_oid:
        return jsonify(error="Invalid product ID"), 400
    product = db().products.find_one({"_id": product_oid})
    if not product:
        return jsonify(error="Product not found"), 404
    dependencies = _product_dependencies(product)
    if any(dependencies.values()):
        return jsonify(error="Product has historical or operational references; archive it instead of deleting it.", dependencies=dependencies), 409
    result = db().products.delete_one({"_id": product_oid})
    if not result.deleted_count:
        return jsonify(error="Product changed; reload before deleting"), 409
    activity(g.user, "delete", "product", product_oid)
    return jsonify(ok=True)


def _collection_for_name(name):
    collection = db().collections.find_one({"name": name, "active": True})
    if not collection:
        raise ValueError("A valid active collection is required")
    code = collection.get("code") or collection_code(collection["name"])
    conflict = db().collections.find_one({"code": code, "_id": {"$ne": collection["_id"]}})
    if conflict:
        raise ValueError(f"Collection code {code} is already used by another collection")
    if collection.get("code") != code:
        db().collections.update_one({"_id": collection["_id"]}, {"$set": {"code": code}})
        collection["code"] = code
    return collection


def _normalize_linesheet_items(raw_items, collection, persist_configurations=True):
    items, seen = [], set()
    for raw in raw_items or []:
        product_code = str(raw.get("product_code") or raw.get("vendor_code") or raw.get("sku") or "").strip()
        source_vendor_code = str(raw.get("source_vendor_code") or raw.get("vendor_code") or product_code).strip()
        source_client_sku = str(raw.get("source_client_sku") or raw.get("client_source_sku") or raw.get("sku") or "").strip()
        description = str(raw.get("description") or raw.get("product_name") or "").strip()
        color = str(raw.get("color") or "").strip()
        set_of = int(raw.get("set_of") or raw.get("component_count") or 1)
        mrp = float(raw.get("mrp") or raw.get("unit_price") or 0)
        currency = str(raw.get("currency") or "INR").upper()
        tax_inclusive = raw.get("tax_inclusive", True)
        sizes = raw.get("sizes") or [key for key, value in (raw.get("size_quantities") or {}).items() if int(value or 0) >= 0]
        sizes = list(dict.fromkeys(str(size).upper() for size in sizes if str(size).upper() in ALLOWED_SIZES))
        if not product_code or not color or not description:
            raise ValueError("Each product row requires product code, description and color")
        if not sizes:
            raise ValueError("Each product row requires at least one size")
        if mrp < 0:
            raise ValueError("MRP cannot be negative")
        if currency not in SUPPORTED_CURRENCIES:
            raise ValueError("Unsupported currency")
        if not isinstance(tax_inclusive, bool):
            raise ValueError("Tax-inclusive state must be boolean")
        sku = linesheet_sku(collection["name"], product_code, color, set_of, collection.get("code"))
        if sku in seen:
            raise ValueError(f"Duplicate product configuration: {sku}")
        seen.add(sku)
        size_quantities = {size: int((raw.get("size_quantities") or {}).get(size, 0)) for size in sizes}
        product_id = raw.get("product_id")
        configuration_id = raw.get("configuration_id")
        if persist_configurations:
            product = db().products.find_one({"product_code": product_code, "collection_id": collection["_id"]})
            if not product:
                internal_product_key = f"PRODUCT-{collection['code']}-{re.sub(r'[^A-Z0-9]+', '-', product_code.upper()).strip('-')}"
                result = db().products.insert_one({"sku": internal_product_key, "product_code": product_code, "name": description, "description": description, "collection": collection["name"], "collection_id": collection["_id"], "base_currency": currency, "tax_inclusive": tax_inclusive, "images": raw.get("images", []), "active": True, "created_at": now(), "updated_at": now()})
                product_id = result.inserted_id
            else:
                product_id = product["_id"]
            configuration = db().product_configurations.find_one_and_update({"linesheet_sku": sku}, {"$setOnInsert": {"product_id": product_id, "collection_id": collection["_id"], "product_code": product_code, "color": color, "set_of": set_of, "linesheet_sku": sku, "created_at": now()}, "$set": {"description": description, "mrp": money(mrp), "currency": currency, "tax_inclusive": tax_inclusive, "source_vendor_code": source_vendor_code, "source_client_sku": source_client_sku, "source_order_quantities": raw.get("size_quantities") or {}, "measurements": raw.get("measurements") or raw.get("size_measurements") or {}, "category": raw.get("category"), "po_reference": raw.get("po_reference"), "delivery_date": raw.get("po_delivery_date"), "images": raw.get("images", []), "updated_at": now()}}, upsert=True, return_document=ReturnDocument.AFTER)
            configuration_id = configuration["_id"]
        items.append({"product_id": product_id, "configuration_id": configuration_id, "product_code": product_code, "vendor_code": source_vendor_code, "source_vendor_code": source_vendor_code, "source_client_sku": source_client_sku, "description": description, "product_name": description, "category": raw.get("category"), "color": color, "set_of": set_of, "sizes": sizes, "size_quantities": size_quantities, "source_order_quantities": size_quantities, "mrp": money(mrp), "currency": currency, "tax_inclusive": tax_inclusive, "measurements": raw.get("measurements") or raw.get("size_measurements") or {}, "po_reference": raw.get("po_reference"), "po_delivery_date": raw.get("po_delivery_date"), "linesheet_sku": sku, "sku": sku, "images": raw.get("images", []), "total_quantity": 0, "total_price": money(0)})
    return items


@api.route("/linesheets", methods=["GET", "POST"])
@auth_required
def linesheets():
    if request.method == "GET":
        return list_response(db().linesheets, search_query(["name", "linesheet_number", "collection", "status", "items.product_code", "items.linesheet_sku"]))
    if g.user["role"] not in {"admin", "sales"}:
        return jsonify(error="Permission denied"), 403
    try:
        data = body(("name", "collection"))
        collection = _collection_for_name(data["collection"])
        items = _normalize_linesheet_items(data.get("items", []), collection)
        doc = {"name": data["name"].strip(), "collection": collection["name"], "collection_id": collection["_id"], "collection_code": collection["code"], "linesheet_number": next_number(db(), "linesheet", "LS"), "items": items, "total_quantity": 0, "total_value": money(0), "status": "draft", "workdrive_status": "not_uploaded", "created_by": g.user["_id"], "created_at": now(), "updated_at": now()}
        result = db().linesheets.insert_one(doc)
        activity(g.user, "create", "linesheet", result.inserted_id)
        return jsonify(id=str(result.inserted_id), linesheet_number=doc["linesheet_number"]), 201
    except (ValueError, DuplicateKeyError) as exc:
        return json_error(exc)


@api.post("/linesheets/<linesheet_id>/convert")
@permission_required("orders:write")
def convert_linesheet(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    if sheet["status"] not in {"confirmed", "shared"}:
        return jsonify(error="Only shared or confirmed linesheets can be converted"), 409
    existing = db().orders.find_one({"source_linesheet_id": sheet["_id"]})
    if existing:
        return jsonify(id=str(existing["_id"]), po_number=existing["po_number"], existing=True)
    order_type = sheet["type"]
    doc = {"po_number": next_number(db(), "po", "RK"), "client_id": sheet["client_id"], "client_name": sheet["client_name"], "order_type": order_type, "source_linesheet_id": sheet["_id"], "items": sheet["items"], "total_quantity": sheet["total_quantity"], "total_amount": sheet["total_value"], "status": "draft", "production_status": "not_started", "po_date": now(), "created_by": g.user["_id"], "created_at": now(), "updated_at": now()}
    result = db().orders.insert_one(doc)
    db().linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"status": "converted_to_order", "order_id": result.inserted_id, "updated_at": now()}})
    db().documents.update_many({"linesheet_id": sheet["_id"], "kind": "purchase_order", "order_id": None}, {"$set": {"order_id": result.inserted_id}})
    activity(g.user, "convert", "linesheet", linesheet_id, {"order_id": str(result.inserted_id)})
    return jsonify(id=str(result.inserted_id), po_number=doc["po_number"]), 201


@api.route("/orders", methods=["GET", "POST"])
@auth_required
def orders():
    if request.method == "GET":
        query = search_query(["po_number", "client_name", "order_type", "status", "production_status", "collection"])
        return list_response(db().orders, query)
    if g.user["role"] not in {"admin", "sales"}:
        return jsonify(error="Permission denied"), 403
    try:
        data = body(("client_id", "order_type", "items"))
        client = db().clients.find_one({"_id": oid(data["client_id"])})
        if not client:
            raise ValueError("Client not found")
        if client.get("status") != "active":
            raise ValueError("Inactive clients cannot be used for new orders")
        if data["order_type"] not in {"mds_outright", "mds_consignment", "direct"}:
            raise ValueError("Invalid order type")
        po_number = data.get("po_number") or next_number(db(), "po", "RK")
        items, total_qty, total_amount = [], 0, 0.0
        for index, item in enumerate(data["items"]):
            qty = int(item.get("quantity", 0))
            if qty <= 0:
                raise ValueError("Order item quantity must be positive")
            price = money(item.get("unit_price"))
            items.append({**item, "line_id": item.get("line_id", str(uuid4())), "quantity": qty, "completed_quantity": 0, "dispatched_quantity": 0, "unit_price": price, "line_total": money(qty * float(price))})
            total_qty += qty
            total_amount += qty * float(price)
        doc = {**data, "po_number": po_number.strip().upper(), "client_id": client["_id"], "client_name": client["name"], "items": items, "total_quantity": total_qty, "total_amount": money(total_amount), "status": data.get("status", "po_received"), "production_status": "not_started", "created_by": g.user["_id"], "created_at": now(), "updated_at": now()}
        result = db().orders.insert_one(doc)
        activity(g.user, "create", "order", result.inserted_id)
        return jsonify(id=str(result.inserted_id), po_number=doc["po_number"]), 201
    except Exception as exc:
        return json_error(exc)


@api.get("/orders/<order_id>")
@auth_required
def order_detail(order_id):
    order = db().orders.find_one({"_id": oid(order_id)})
    if not order:
        return jsonify(error="Order not found"), 404
    payload = {
        "order": order,
        "client": db().clients.find_one({"_id": order["client_id"]}),
        "documents": list(db().documents.find({"order_id": order["_id"]}).sort("version", DESCENDING)),
        "production": list(db().production_batches.find({"order_id": order["_id"]})),
        "stock_ledger": list(db().stock_ledger.find({"related_po": order["po_number"]}).sort("created_at", DESCENDING)),
        "activity": list(db().activity_log.find({"entity_id": order_id}).sort("created_at", DESCENDING)),
    }
    return jsonify(serialize(payload))


@api.post("/orders/<order_id>/confirm")
@permission_required("orders:write")
def confirm_order(order_id):
    order = db().orders.find_one({"_id": oid(order_id)})
    if not order:
        return jsonify(error="Order not found"), 404
    if order["status"] not in {"draft", "po_received"}:
        return jsonify(error="Order cannot be confirmed from its current status"), 409
    db().orders.update_one({"_id": order["_id"]}, {"$set": {"status": "confirmed", "confirmed_at": now(), "updated_at": now()}})
    activity(g.user, "confirm", "order", order_id)
    return jsonify(status="confirmed")


@api.post("/orders/<order_id>/production")
@permission_required("production:write")
def create_production(order_id):
    order = db().orders.find_one({"_id": oid(order_id)})
    if not order:
        return jsonify(error="Order not found"), 404
    try:
        data = body(("line_id", "quantity"))
        line = next((item for item in order["items"] if item["line_id"] == data["line_id"]), None)
        if not line:
            raise ValueError("Order line not found")
        allocated = sum(item.get("quantity", 0) for item in db().production_batches.find({"order_id": order["_id"], "line_id": data["line_id"]}))
        quantity = int(data["quantity"])
        if quantity <= 0 or allocated + quantity > line["quantity"]:
            raise ValueError("Production allocation exceeds outstanding order quantity")
        doc = {"order_id": order["_id"], "po_number": order["po_number"], "client_name": order["client_name"], "order_type": order["order_type"], "line_id": line["line_id"], "sku": line["sku"], "product_name": line.get("product_name"), "color": line.get("color"), "size": line.get("size"), "quantity": quantity, "stage_quantities": {"designing": quantity}, "current_stage": "designing", "completed_quantity": 0, "revision": 0, "expected_completion_date": data.get("expected_completion_date"), "created_at": now(), "updated_at": now()}
        result = db().production_batches.insert_one(doc)
        db().orders.update_one({"_id": order["_id"]}, {"$set": {"status": "in_production", "production_status": "designing", "updated_at": now()}})
        activity(g.user, "create_production", "order", order_id, {"batch_id": str(result.inserted_id), "quantity": quantity})
        return jsonify(id=str(result.inserted_id)), 201
    except Exception as exc:
        return json_error(exc)


@api.get("/production")
@auth_required
def production():
    query = search_query(["po_number", "client_name", "order_type", "sku", "product_name", "current_stage"])
    if request.args.get("delayed") == "true":
        query["expected_completion_date"] = {"$lt": now().date().isoformat()}
        query["current_stage"] = {"$ne": "dispatched"}
    return list_response(db().production_batches, query)


@api.post("/production/<batch_id>/move")
@permission_required("production:write")
def production_move(batch_id):
    batch = db().production_batches.find_one({"_id": oid(batch_id)})
    if not batch:
        return jsonify(error="Production batch not found"), 404
    try:
        data = body(("to_stage", "quantity"))
        updated, movement = move_production(batch, data.get("from_stage", batch.get("current_stage")), data["to_stage"], data["quantity"], g.user, data.get("notes"), data.get("rejected", 0), data.get("rework", 0))
        activity(g.user, "move_stage", "production_batch", batch_id, {"from": movement["from_stage"], "to": movement["to_stage"], "quantity": movement["quantity"]})
        return jsonify(batch=serialize(updated), movement=serialize(movement))
    except Exception as exc:
        return json_error(exc, 409 if isinstance(exc, RuntimeError) else 400)


@api.route("/stock", methods=["GET", "POST"])
@auth_required
def stock():
    if request.method == "GET":
        query = search_query(["sku", "color", "size", "location"])
        if request.args.get("low"):
            query["available"] = {"$lte": int(request.args["low"])}
        return list_response(db().stock_balances, query, ("sku", 1))
    if g.user["role"] not in {"admin", "production_inventory"}:
        return jsonify(error="Permission denied"), 403
    try:
        data = body(("sku", "color", "size", "quantity", "transaction_type"))
        if data.get("location") and data["location"] not in ("main", *INVENTORY_LOCATIONS):
            raise ValueError("Invalid inventory location")
        ledger, created = mutate_stock(**{k: data.get(k) for k in ("sku", "color", "size", "quantity", "transaction_type", "location", "client_id", "po_number", "notes", "idempotency_key", "allow_negative") if data.get(k) is not None}, user=g.user)
        activity(g.user, data["transaction_type"], "stock", ledger["transaction_id"], {"sku": data["sku"], "quantity": data["quantity"]})
        return jsonify(transaction=serialize(ledger), created=created), 201 if created else 200
    except Exception as exc:
        return json_error(exc, 409 if isinstance(exc, RuntimeError) else 400)


@api.get("/stock/ledger")
@auth_required
def stock_ledger():
    return list_response(db().stock_ledger, search_query(["transaction_id", "sku", "color", "size", "transaction_type", "related_po", "user_email"]))


@api.get("/inventory/locations")
@auth_required
def inventory_locations():
    return jsonify(locations=[{"name": name, "available_for_sale": name not in NON_SELLABLE_LOCATIONS} for name in INVENTORY_LOCATIONS])


INVENTORY_ADJUSTMENT_REASONS = {
    "production_received": "production_completion",
    "manual_adjustment": "adjustment",
    "damage": "damaged",
    "return": "return",
    "transfer": "adjustment",
    "consignment": "consignment_sent",
    "other": "adjustment",
}


def _product_inventory_rows(product):
    variants = list(db().inventory_variants.find({"product_id": product["_id"]}).sort("size", 1))
    rows = []
    for variant in variants:
        balances = list(db().stock_balances.find({"sku": variant["inventory_sku"], "color": variant.get("color"), "size": variant["size"]}))
        physical = sum(int(item.get("physical", 0)) for item in balances)
        reserved = sum(int(item.get("reserved", 0)) for item in balances)
        rows.append({**variant, "physical": physical, "reserved": reserved, "available": physical - reserved, "total": physical, "locations": serialize(balances)})
    return rows


@api.route("/products/<product_id>/inventory", methods=["GET"])
@auth_required
def product_inventory(product_id):
    product_oid = oid(product_id)
    if not product_oid:
        return jsonify(error="Invalid product ID"), 400
    product = db().products.find_one({"_id": product_oid})
    if not product:
        return jsonify(error="Product not found"), 404
    return jsonify(items=serialize(_product_inventory_rows(product)))


@api.post("/products/<product_id>/inventory/adjust")
@permission_required("stock:write")
def adjust_product_inventory(product_id):
    product_oid = oid(product_id)
    if not product_oid:
        return jsonify(error="Invalid product ID"), 400
    product = db().products.find_one({"_id": product_oid})
    if not product:
        return jsonify(error="Product not found"), 404
    try:
        data = body(("variant_id", "quantity", "direction", "reason"))
        variant = db().inventory_variants.find_one({"_id": oid(data["variant_id"]), "product_id": product_oid})
        if not variant:
            return jsonify(error="Variant not found for this product"), 404
        quantity = int(data["quantity"])
        if quantity <= 0:
            raise ValueError("Quantity must be greater than zero")
        direction = str(data["direction"]).strip().lower()
        if direction not in {"add", "remove"}:
            raise ValueError("Direction must be add or remove")
        reason = str(data["reason"]).strip().lower()
        transaction_type = INVENTORY_ADJUSTMENT_REASONS.get(reason)
        if not transaction_type:
            raise ValueError("Invalid inventory adjustment reason")
        signed = quantity if direction == "add" else -quantity
        request_id = data.get("idempotency_key") or request.headers.get("Idempotency-Key")
        if not request_id:
            raise ValueError("An idempotency key is required")
        ledger, created = mutate_stock(sku=variant["inventory_sku"], color=variant.get("color"), size=variant["size"], quantity=signed, transaction_type=transaction_type, user=g.user, location=data.get("location") or "main", notes=data.get("notes"), idempotency_key=request_id)
        activity(g.user, "stock_adjustment", "inventory_variant", variant["_id"], {"reason": reason, "quantity": signed, "transaction_id": ledger["transaction_id"]})
        return jsonify(transaction=serialize(ledger), created=created, inventory=serialize(_product_inventory_rows(product))), 201 if created else 200
    except Exception as exc:
        return json_error(exc, 409 if isinstance(exc, (RuntimeError, DuplicateKeyError)) else 400)


@api.get("/pricing/convert")
@auth_required
def convert_price():
    try:
        amount = request.args.get("amount")
        target = request.args.get("to", "INR")
        provider = FxProvider(current_app.config.get("FX_API_URL"), current_app.config.get("FX_API_KEY"), current_app.config.get("FX_API_TIMEOUT_SECONDS", 5))
        converted = provider.convert(amount, target)
        return jsonify(amount=str(amount), base_currency="INR", target_currency=target.upper(), converted_amount=str(converted), authoritative=False)
    except FxProviderError as exc:
        return jsonify(error=str(exc)), 503


@api.route("/inventory/variants", methods=["GET", "POST"])
@auth_required
def inventory_variants():
    if request.method == "GET":
        return list_response(db().inventory_variants, search_query(["inventory_sku", "linesheet_sku", "product_code", "color", "size"]), ("inventory_sku", 1))
    if g.user["role"] not in {"admin", "production_inventory"}:
        return jsonify(error="Permission denied"), 403
    try:
        data = body(("linesheet_sku", "sizes"))
        configuration = db().product_configurations.find_one({"linesheet_sku": data["linesheet_sku"]})
        if not configuration:
            raise ValueError("Product configuration was not found")
        created, reused = [], []
        quantities = data.get("quantities") or {}
        for size in dict.fromkeys(data["sizes"]):
            sku = inventory_sku(configuration["linesheet_sku"], size)
            existing = db().inventory_variants.find_one({"inventory_sku": sku})
            if existing:
                reused.append(sku)
            else:
                db().inventory_variants.insert_one({"configuration_id": configuration["_id"], "product_id": configuration["product_id"], "linesheet_sku": configuration["linesheet_sku"], "inventory_sku": sku, "product_code": configuration["product_code"], "color": configuration["color"], "set_of": configuration["set_of"], "size": size, "measurements": (configuration.get("measurements") or {}).get(size, {}), "source_client_sku": configuration.get("source_client_sku"), "source_vendor_code": configuration.get("source_vendor_code"), "source_order_quantity": int((configuration.get("source_order_quantities") or {}).get(size, 0)), "created_at": now(), "updated_at": now()})
                created.append(sku)
            quantity = int(quantities.get(size, 0))
            if quantity < 0:
                raise ValueError("Inventory quantity cannot be negative")
            if quantity > 0:
                location = data.get("location")
                if not location:
                    raise ValueError("A physical stock location is required when opening quantity is supplied")
                if location not in INVENTORY_LOCATIONS:
                    raise ValueError("Invalid inventory location")
                mutate_stock(sku=sku, color=configuration["color"], size=size, quantity=quantity, transaction_type="opening_stock", user=g.user, location=location, available_for_sale=data.get("available_for_sale"), idempotency_key=f"inventory:{configuration['_id']}:{size}:{data.get('request_id') or uuid4()}")
        return jsonify(created=created, reused=reused), 201
    except (ValueError, DuplicateKeyError) as exc:
        return json_error(exc)


def _product_variant_configuration(product):
    configuration = db().product_configurations.find_one({"product_id": product["_id"]}, sort=[("created_at", 1)])
    if configuration:
        return configuration
    collection_id = (product.get("collection_ids") or [product.get("collection_id")])[0] if (product.get("collection_ids") or product.get("collection_id")) else None
    document = {
        "product_id": product["_id"], "collection_id": collection_id, "product_code": product.get("product_code") or product.get("sku"),
        "color": (product.get("colors") or product.get("colours") or [product.get("color") or ""])[0],
        "set_of": 1, "linesheet_sku": product.get("sku"), "description": product.get("description") or product.get("name"),
        "created_at": now(), "updated_at": now(),
    }
    return db().product_configurations.find_one_and_update({"product_id": product["_id"]}, {"$setOnInsert": document}, upsert=True, return_document=ReturnDocument.AFTER)


def _variant_references(variant):
    sku = variant.get("inventory_sku")
    return {
        "stock_balances": db().stock_balances.count_documents({"sku": sku}),
        "stock_ledger": db().stock_ledger.count_documents({"sku": sku}),
        "orders": db().orders.count_documents({"items.sku": sku}),
        "production": db().production_movements.count_documents({"sku": sku}),
    }


@api.route("/products/<product_id>/variants", methods=["GET", "POST"])
@auth_required
def product_variants(product_id):
    product_oid = oid(product_id)
    if not product_oid:
        return jsonify(error="Invalid product ID"), 400
    product = db().products.find_one({"_id": product_oid})
    if not product:
        return jsonify(error="Product not found"), 404
    if request.method == "GET":
        variants = list(db().inventory_variants.find({"product_id": product_oid}).sort("size", 1))
        return jsonify(items=serialize(variants))
    if "*" not in effective_permissions(g.user) and "settings:write" not in effective_permissions(g.user):
        return jsonify(error="Permission denied"), 403
    try:
        data = body(("size",))
        size = str(data["size"]).strip().upper()
        if size not in ALLOWED_SIZES:
            raise ValueError("Size must be XS, S, M, L or XL")
        configuration = _product_variant_configuration(product)
        if db().inventory_variants.find_one({"product_id": product_oid, "size": size}):
            return jsonify(error="A variant for this product and size already exists"), 409
        inventory_sku = str(data.get("inventory_sku") or f"{product.get('sku')}-{size}").strip().upper()
        if db().inventory_variants.find_one({"inventory_sku": inventory_sku}):
            return jsonify(error="Variant SKU already exists"), 409
        active = data.get("active", True)
        if not isinstance(active, bool):
            raise ValueError("Active state must be boolean")
        variant = {"configuration_id": configuration["_id"], "product_id": product_oid, "linesheet_sku": configuration.get("linesheet_sku"), "inventory_sku": inventory_sku, "product_code": configuration.get("product_code"), "color": configuration.get("color"), "size": size, "active": active, "status": "active" if active else "archived", "metadata": data.get("metadata") or {}, "created_at": now(), "updated_at": now()}
        result = db().inventory_variants.insert_one(variant)
        activity(g.user, "create", "inventory_variant", result.inserted_id, {"product_id": product_id, "size": size})
        return jsonify(serialize(db().inventory_variants.find_one({"_id": result.inserted_id}))), 201
    except (ValueError, DuplicateKeyError) as exc:
        return json_error(exc, 409 if isinstance(exc, DuplicateKeyError) else 400)


@api.route("/products/<product_id>/variants/<variant_id>", methods=["PATCH", "DELETE"])
@permission_required("settings:write")
def product_variant_detail(product_id, variant_id):
    product_oid, variant_oid = oid(product_id), oid(variant_id)
    if not product_oid or not variant_oid:
        return jsonify(error="Invalid product or variant ID"), 400
    variant = db().inventory_variants.find_one({"_id": variant_oid, "product_id": product_oid})
    if not variant:
        return jsonify(error="Variant not found"), 404
    references = _variant_references(variant)
    if request.method == "DELETE":
        if any(references.values()):
            return jsonify(error="Variant has historical stock or order references and cannot be removed", references=references), 409
        db().inventory_variants.delete_one({"_id": variant_oid})
        activity(g.user, "delete", "inventory_variant", variant_oid)
        return jsonify(ok=True)
    data = request.get_json(silent=True) or {}
    allowed = {"size", "inventory_sku", "active", "metadata"}
    if set(data) - allowed:
        return jsonify(error="Only size, inventory_sku, active and metadata can be changed"), 400
    changes = {}
    if "size" in data:
        size = str(data["size"]).strip().upper()
        if size not in ALLOWED_SIZES:
            return jsonify(error="Size must be XS, S, M, L or XL"), 400
        changes["size"] = size
    if "inventory_sku" in data:
        sku = str(data["inventory_sku"]).strip().upper()
        if not sku or db().inventory_variants.find_one({"inventory_sku": sku, "_id": {"$ne": variant_oid}}):
            return jsonify(error="Variant SKU is blank or already exists"), 409
        changes["inventory_sku"] = sku
    if "active" in data:
        if not isinstance(data["active"], bool):
            return jsonify(error="Active state must be boolean"), 400
        changes["active"] = data["active"]
        changes["status"] = "active" if data["active"] else "archived"
    if "metadata" in data:
        if not isinstance(data["metadata"], dict):
            return jsonify(error="Metadata must be an object"), 400
        changes["metadata"] = data["metadata"]
    if not changes:
        return jsonify(error="No editable variant fields supplied"), 400
    if "size" in changes and db().inventory_variants.find_one({"configuration_id": variant["configuration_id"], "size": changes["size"], "_id": {"$ne": variant_oid}}):
        return jsonify(error="A variant for this configuration and size already exists"), 409
    changes["updated_at"] = now()
    db().inventory_variants.update_one({"_id": variant_oid}, {"$set": changes})
    activity(g.user, "update", "inventory_variant", variant_oid, {"fields": sorted(changes)})
    return jsonify(serialize(db().inventory_variants.find_one({"_id": variant_oid})))


@api.post("/orders/<order_id>/reserve")
@permission_required("stock:write")
def reserve_order(order_id):
    order = db().orders.find_one({"_id": oid(order_id)})
    if not order:
        return jsonify(error="Order not found"), 404
    try:
        data = body(("line_id", "quantity"))
        line = next((item for item in order["items"] if item["line_id"] == data["line_id"]), None)
        if not line:
            raise ValueError("Order line not found")
        qty = int(data["quantity"])
        if qty <= 0 or qty > line["quantity"] - int(line.get("reserved_quantity", 0)):
            raise ValueError("Reservation exceeds outstanding line quantity")
        balance = db().stock_balances.find_one({"sku": line["sku"], "color": line.get("color"), "size": line.get("size"), "location": "main", "client_id": None})
        if not balance or int(balance.get("available", balance.get("physical", 0) - balance.get("reserved", 0))) < qty:
            raise ValueError("Insufficient available stock")
        mutate_stock(sku=line["sku"], color=line.get("color"), size=line.get("size"), quantity=qty, transaction_type="reserve", user=g.user, po_number=order["po_number"], idempotency_key=request.headers.get("Idempotency-Key"))
        db().orders.update_one({"_id": order["_id"], "items.line_id": line["line_id"]}, {"$inc": {"items.$.reserved_quantity": qty}, "$set": {"updated_at": now()}})
        return jsonify(reserved=qty)
    except Exception as exc:
        return json_error(exc)


@api.post("/orders/<order_id>/dispatch")
@permission_required("stock:write")
def dispatch_order(order_id):
    order = db().orders.find_one({"_id": oid(order_id)})
    if not order:
        return jsonify(error="Order not found"), 404
    try:
        data = body(("line_id", "quantity"))
        line = next((item for item in order["items"] if item["line_id"] == data["line_id"]), None)
        if not line:
            raise ValueError("Order line not found")
        qty = int(data["quantity"])
        outstanding = line["quantity"] - int(line.get("dispatched_quantity", 0))
        if qty <= 0 or qty > outstanding or qty > int(line.get("reserved_quantity", 0)):
            raise ValueError("Dispatch exceeds reserved or outstanding quantity")
        key = request.headers.get("Idempotency-Key") or data.get("idempotency_key")
        if not key:
            raise ValueError("Idempotency-Key header is required")
        ledger, created = mutate_stock(sku=line["sku"], color=line.get("color"), size=line.get("size"), quantity=-qty, transaction_type="dispatch", user=g.user, po_number=order["po_number"], idempotency_key=key)
        if created:
            mutate_stock(sku=line["sku"], color=line.get("color"), size=line.get("size"), quantity=-qty, transaction_type="release", user=g.user, po_number=order["po_number"], idempotency_key=f"{key}:release")
            db().orders.update_one({"_id": order["_id"], "items.line_id": line["line_id"]}, {"$inc": {"items.$.dispatched_quantity": qty, "items.$.reserved_quantity": -qty}, "$set": {"updated_at": now()}})
            refreshed = db().orders.find_one({"_id": order["_id"]})
            dispatched = sum(i.get("dispatched_quantity", 0) for i in refreshed["items"])
            status = "executed" if dispatched == refreshed["total_quantity"] else "partially_dispatched"
            db().orders.update_one({"_id": order["_id"]}, {"$set": {"status": status, "executed_at": now() if status == "executed" else None}})
            activity(g.user, "dispatch", "order", order_id, {"line_id": line["line_id"], "quantity": qty})
        return jsonify(dispatched=qty, created=created)
    except Exception as exc:
        return json_error(exc)


@api.route("/orders/<order_id>/documents", methods=["GET", "POST"])
@auth_required
def order_documents(order_id):
    order = db().orders.find_one({"_id": oid(order_id)})
    if not order:
        return jsonify(error="Order not found"), 404
    if request.method == "GET":
        return jsonify(items=serialize(list(db().documents.find({"order_id": order["_id"]}).sort("version", DESCENDING))))
    if g.user["role"] not in {"admin", "sales"}:
        return jsonify(error="Permission denied"), 403
    uploaded = request.files.get("file")
    if not uploaded or uploaded.mimetype != "application/pdf":
        return jsonify(error="A PDF file is required"), 400
    data = uploaded.read()
    if not data.startswith(b"%PDF"):
        return jsonify(error="Uploaded file is not a valid PDF"), 400
    family_id = request.form.get("family_id") or str(uuid4())
    previous = db().documents.find_one({"family_id": family_id}, sort=[("version", DESCENDING)])
    version = int(previous["version"]) + 1 if previous else 1
    folder = Path(current_app.config["UPLOAD_DIR"]) / str(order["_id"])
    folder.mkdir(parents=True, exist_ok=True)
    stored_name = f"{family_id}-v{version}.pdf"
    (folder / stored_name).write_bytes(data)
    doc = {"family_id": family_id, "version": version, "order_id": order["_id"], "client_id": order["client_id"], "kind": "purchase_order", "original_name": secure_filename(uploaded.filename or "purchase-order.pdf"), "stored_path": str(folder / stored_name), "content_type": "application/pdf", "size": len(data), "sha256": hashlib.sha256(data).hexdigest(), "uploaded_by": g.user["_id"], "uploaded_by_email": g.user["email"], "created_at": now()}
    result = db().documents.insert_one(doc)
    activity(g.user, "upload_document", "order", order_id, {"document_id": str(result.inserted_id), "version": version})
    return jsonify(id=str(result.inserted_id), family_id=family_id, version=version), 201


@api.get("/documents/<document_id>/download")
@auth_required
def download_document(document_id):
    doc = db().documents.find_one({"_id": oid(document_id)})
    if not doc or not os.path.exists(doc["stored_path"]):
        return jsonify(error="Document not found"), 404
    return send_file(doc["stored_path"], mimetype=doc["content_type"], as_attachment=request.args.get("preview") != "true", download_name=doc["original_name"])


@api.get("/dashboard")
@auth_required
def dashboard():
    period = request.args.get("period", "month")
    start = {"today": now() - timedelta(days=1), "week": now() - timedelta(days=7), "month": now() - timedelta(days=30)}.get(period)
    base = {"created_at": {"$gte": start}} if start else {}
    sections = {}
    for key, order_type in (("mds_outright", "mds_outright"), ("mds_consignment", "mds_consignment"), ("direct", "direct")):
        query = {**base, "order_type": order_type}
        sections[key] = {
            "received": db().orders.count_documents(query),
            "executed": db().orders.count_documents({**query, "status": "executed"}),
            "pending": db().orders.count_documents({**query, "status": {"$nin": ["executed", "cancelled"]}}),
            "ready_to_dispatch": db().orders.count_documents({**query, "status": "ready_to_dispatch"}),
            "outstanding_quantity": sum(max(0, item.get("total_quantity", 0) - sum(line.get("dispatched_quantity", 0) for line in item.get("items", []))) for item in db().orders.find(query, {"total_quantity": 1, "items.dispatched_quantity": 1})),
        }
    consignment = list(db().stock_ledger.aggregate([{"$match": {"transaction_type": {"$in": ["consignment_sent", "consignment_sold", "consignment_returned"]}}}, {"$group": {"_id": "$transaction_type", "quantity": {"$sum": "$quantity"}}}]))
    sections["mds_consignment"]["stock"] = {item["_id"]: item["quantity"] for item in consignment}
    production_metrics = {stage: 0 for stage in STAGES}
    for item in db().production_batches.aggregate([{"$project": {"entries": {"$objectToArray": "$stage_quantities"}}}, {"$unwind": "$entries"}, {"$group": {"_id": "$entries.k", "quantity": {"$sum": "$entries.v"}}}]):
        production_metrics[item["_id"]] = item["quantity"]
    return jsonify(sections=sections, production=production_metrics, generated_at=now().isoformat())


@api.get("/reports/<report_name>.xlsx")
@auth_required
def report(report_name):
    allowed = {
        "orders": (db().orders, ["po_number", "client_name", "order_type", "status", "total_quantity", "total_amount"]),
        "production": (db().production_batches, ["po_number", "client_name", "sku", "color", "size", "quantity", "current_stage", "expected_completion_date"]),
        "stock": (db().stock_balances, ["sku", "color", "size", "location", "physical", "reserved", "available", "consignment"]),
        "stock-ledger": (db().stock_ledger, ["transaction_id", "sku", "color", "size", "quantity", "transaction_type", "related_po", "user_email", "created_at"]),
    }
    if report_name not in allowed:
        return jsonify(error="Unknown report"), 404
    collection, columns = allowed[report_name]
    records = collection.find(search_query(columns))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report_name[:31]
    sheet.append(columns)
    for record in records:
        sheet.append([str(serialize(record.get(column, ""))) for column in columns])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"rk-{report_name}-{now().date().isoformat()}.xlsx")


@api.get("/settings")
@auth_required
def settings():
    configured = db().settings.find_one({"_id": "global"}) or {}
    return jsonify(serialize({"sizes": configured.get("sizes", ["XS", "S", "M", "L", "XL", "XXL"]), "production_stages": configured.get("production_stages", STAGES), "categories": configured.get("categories", [])}))


CATEGORY_NAME_MAX_LENGTH = 80
CATEGORY_REVISION_FIELD = "categories_revision"
COLLECTION_REVISION_FIELD = "collections_revision"


def _category_state():
    configured = db().settings.find_one({"_id": "global"}, {"categories": 1, CATEGORY_REVISION_FIELD: 1}) or {}
    values = configured.get("categories")
    issues = []
    categories = []
    if "categories" not in configured:
        values = []
    if not isinstance(values, list):
        issues.append({"kind": "malformed", "index": None, "value": values})
        values = []
    seen = set()
    for index, value in enumerate(values):
        if not isinstance(value, str):
            issues.append({"kind": "malformed", "index": index, "value": value})
            continue
        category = value.strip()
        if not category:
            issues.append({"kind": "blank", "index": index, "value": value})
            continue
        key = category.casefold()
        if key in seen:
            issues.append({"kind": "duplicate", "index": index, "value": value})
            continue
        seen.add(key)
        categories.append(category)
    revision = configured.get(CATEGORY_REVISION_FIELD, 0)
    if not isinstance(revision, int) or revision < 0:
        issues.append({"kind": "malformed_revision", "index": None, "value": revision})
        revision = 0
    return {"categories": categories, "revision": revision, "issues": issues}


def _category_usage(categories):
    return {
        category: db().products.count_documents({
            "category": {"$regex": rf"^\s*{re.escape(category)}\s*$", "$options": "i"},
        })
        for category in categories
    }


def _category_response(state):
    return {
        "categories": state["categories"],
        "usage": _category_usage(state["categories"]),
        "revision": state["revision"],
        "legacy_issues": state["issues"],
    }


def _collection_state():
    records = list(db().collections.find({}).sort("position", 1))
    issues = []
    valid = []
    seen = set()
    for record in records:
        name = record.get("name")
        if not isinstance(name, str) or not name.strip():
            issues.append({"kind": "malformed", "id": str(record.get("_id")), "value": name})
            continue
        clean = name.strip()
        key = clean.casefold()
        if key in seen:
            issues.append({"kind": "duplicate", "id": str(record.get("_id")), "value": name})
            continue
        seen.add(key)
        valid.append(record)
    setting = db().settings.find_one({"_id": "global"}, {COLLECTION_REVISION_FIELD: 1, "collections_revision_recovery_required": 1}) or {}
    if setting.get("collections_revision_recovery_required"):
        issues.append({"kind": "revision_recovery_required", "value": True})
    revision = setting.get(COLLECTION_REVISION_FIELD, 0)
    if not isinstance(revision, int) or revision < 0:
        issues.append({"kind": "malformed_revision", "value": revision})
        revision = 0
    return {"records": valid, "revision": revision, "issues": issues}


def _collection_usage(record):
    name = record.get("name")
    identifier = record.get("_id")
    queries = [{"collection_id": identifier}, {"collection": name}]
    return sum(collection.count_documents({"$or": queries}) for collection in (db().products, db().product_configurations, db().linesheets, db().orders, db().documents))


def _collection_response(state):
    return {"collections": [{**serialize(record), "usage": _collection_usage(record)} for record in state["records"]], "revision": state["revision"], "legacy_issues": state["issues"]}


def _advance_collection_revision(state):
    if not db().settings.find_one({"_id": "global"}, {"_id": 1}):
        try:
            db().settings.insert_one({"_id": "global", COLLECTION_REVISION_FIELD: 0})
        except DuplicateKeyError:
            pass
    result = db().settings.update_one(
        {"_id": "global", "$or": [{COLLECTION_REVISION_FIELD: state["revision"]}, {COLLECTION_REVISION_FIELD: {"$exists": False}}]},
        {"$set": {COLLECTION_REVISION_FIELD: state["revision"] + 1, "updated_at": now(), "updated_by": g.user["_id"]}},
        upsert=False,
    )
    if result.matched_count:
        return True
    db().settings.update_one({"_id": "global"}, {"$set": {"collections_revision_recovery_required": True, "updated_at": now()}})
    return False


def _validated_category_payload(payload):
    if not isinstance(payload, dict) or set(payload) - {"categories", "revision"} or "categories" not in payload:
        raise ValueError("Only the categories and revision fields are accepted")
    categories = payload.get("categories")
    if not isinstance(categories, list) or any(not isinstance(value, str) for value in categories):
        raise ValueError("categories must be an array of strings")

    normalized = []
    seen = set()
    for value in categories:
        category = value.strip()
        if not category:
            raise ValueError("Category names cannot be empty")
        if len(category) > CATEGORY_NAME_MAX_LENGTH:
            raise ValueError(f"Category names must be {CATEGORY_NAME_MAX_LENGTH} characters or fewer")
        key = category.casefold()
        if key in seen:
            raise ValueError("Category names must be unique, ignoring capitalization")
        seen.add(key)
        normalized.append(category)
    return normalized


@api.get("/settings/categories")
@auth_required
def categories():
    # A missing settings document is a valid empty state. This read must not
    # create global settings as a side effect.
    return jsonify(serialize(_category_response(_category_state())))


@api.put("/settings/categories")
@permission_required("settings:write")
def update_categories():
    try:
        requested = _validated_category_payload(request.get_json(silent=True))
        state = _category_state()
        if state["issues"]:
            return jsonify(error="Existing category data requires administrative cleanup before it can be changed.", legacy_issues=state["issues"]), 409
        requested_revision = request.get_json(silent=True).get("revision") if isinstance(request.get_json(silent=True), dict) else None
        if not isinstance(requested_revision, int) or requested_revision != state["revision"]:
            return jsonify(error="Category data changed; reload before saving.", current_revision=state["revision"]), 409
        existing = state["categories"]
        requested_keys = {value.casefold() for value in requested}
        removed = [value for value in existing if value.casefold() not in requested_keys]
        usage = _category_usage(removed)
        in_use = [value for value in removed if usage.get(value, 0) > 0]
        if in_use:
            names = ", ".join(in_use)
            return jsonify(error=f"Cannot remove or rename category '{names}' because products use it.", in_use=in_use), 409

        timestamp = now()
        next_revision = state["revision"] + 1
        result = db().settings.update_one(
            {"_id": "global", "$or": [{CATEGORY_REVISION_FIELD: state["revision"]}, {CATEGORY_REVISION_FIELD: {"$exists": False}}]},
            {"$set": {"categories": requested, CATEGORY_REVISION_FIELD: next_revision, "updated_at": timestamp, "updated_by": g.user["_id"]}},
            upsert=False,
        )
        if not result.matched_count:
            if db().settings.find_one({"_id": "global"}, {"_id": 1}):
                return jsonify(error="Category data changed; reload before saving.", current_revision=_category_state()["revision"]), 409
            try:
                db().settings.insert_one({"_id": "global", "categories": requested, CATEGORY_REVISION_FIELD: next_revision, "updated_at": timestamp, "updated_by": g.user["_id"]})
            except DuplicateKeyError:
                return jsonify(error="Category data changed; reload before saving.", current_revision=_category_state()["revision"]), 409
        activity(g.user, "update", "settings", "global", {"fields": ["categories"], "category_count": len(requested)})
        return jsonify(ok=True, **_category_response(_category_state()))
    except ValueError as exc:
        return json_error(exc)


@api.patch("/settings")
@permission_required("settings:write")
def update_settings():
    changes = body()
    db().settings.update_one({"_id": "global"}, {"$set": {**changes, "updated_at": now(), "updated_by": g.user["_id"]}}, upsert=True)
    activity(g.user, "update", "settings", "global", changes)
    return jsonify(ok=True)


@api.route("/collections", methods=["GET", "POST"])
@auth_required
def collections():
    if request.method == "GET":
        query = {} if request.args.get("include_archived") == "true" else {"active": True}
        state = _collection_state()
        items = [record for record in state["records"] if not query or query.items() <= record.items()]
        return jsonify(items=serialize(items), revision=state["revision"], legacy_issues=state["issues"])
    if "*" not in effective_permissions(g.user) and "settings:write" not in effective_permissions(g.user):
        return jsonify(error="Permission denied"), 403
    try:
        data = body(("name",))
        state = _collection_state()
        if state["issues"]:
            return jsonify(error="Legacy collection data must be cleaned up before changes are allowed", legacy_issues=state["issues"]), 409
        if data.get("revision") != state["revision"]:
            return jsonify(error="Collection configuration changed; reload before saving", revision=state["revision"]), 409
        if not isinstance(data["name"], str) or not data["name"].strip() or db().collections.find_one({"name": {"$regex": f"^{re.escape(data['name'].strip())}$", "$options": "i"}}):
            return jsonify(error="Collection name is blank or already in use"), 400
        slug = "-".join(data["name"].strip().lower().split())
        position = data.get("position", db().collections.count_documents({}) + 1)
        code = collection_code(data["name"], data.get("code"))
        if db().collections.find_one({"code": code}):
            raise ValueError(f"Collection code {code} is already in use")
        result = db().collections.insert_one({"name": data["name"].strip(), "slug": slug, "code": code, "position": int(position), "active": True, "created_at": now()})
        if not _advance_collection_revision(state):
            return jsonify(error="Collection changed during save; recovery is required before further edits"), 503
        activity(g.user, "create", "collection", result.inserted_id)
        return jsonify(id=str(result.inserted_id)), 201
    except Exception as exc:
        return json_error(exc)


@api.patch("/collections/<collection_id>")
@permission_required("settings:write")
def collection_update(collection_id):
    changes = body()
    state = _collection_state()
    if state["issues"]:
        return jsonify(error="Legacy collection data must be cleaned up before changes are allowed", legacy_issues=state["issues"]), 409
    if changes.get("revision") != state["revision"]:
        return jsonify(error="Collection configuration changed; reload before saving", revision=state["revision"]), 409
    current = db().collections.find_one({"_id": oid(collection_id)})
    if not current:
        return jsonify(error="Collection not found"), 404
    allowed = {key: value for key, value in changes.items() if key in {"name", "code", "position", "active"}}
    if "name" in allowed:
        if _collection_usage(current):
            return jsonify(error="Collection is in use and cannot be renamed"), 409
        if not isinstance(allowed["name"], str) or not allowed["name"].strip():
            return jsonify(error="Collection name cannot be empty"), 400
        allowed["slug"] = "-".join(allowed["name"].strip().lower().split())
    if "code" in allowed:
        allowed["code"] = collection_code(allowed.get("name") or "Collection", allowed["code"])
        if db().collections.find_one({"code": allowed["code"], "_id": {"$ne": oid(collection_id)}}):
            return jsonify(error="Collection code is already in use"), 400
    allowed["updated_at"] = now()
    result = db().collections.update_one({"_id": oid(collection_id)}, {"$set": allowed})
    if not result.matched_count:
        return jsonify(error="Collection not found"), 404
    activity(g.user, "update", "collection", collection_id, allowed)
    if not _advance_collection_revision(state):
        return jsonify(error="Collection changed during save; recovery is required before further edits"), 503
    return jsonify(ok=True)


@api.delete("/collections/<collection_id>")
@permission_required("settings:write")
def collection_delete(collection_id):
    state = _collection_state()
    if state["issues"]:
        return jsonify(error="Legacy collection data must be cleaned up before changes are allowed", legacy_issues=state["issues"]), 409
    if request.args.get("revision", type=int) != state["revision"]:
        return jsonify(error="Collection configuration changed; reload before saving", revision=state["revision"]), 409
    current = db().collections.find_one({"_id": oid(collection_id)})
    if not current:
        return jsonify(error="Collection not found"), 404
    if _collection_usage(current):
        return jsonify(error="Collection is in use and cannot be removed"), 409
    result = db().collections.delete_one({"_id": oid(collection_id)})
    if not result.deleted_count:
        return jsonify(error="Collection changed; reload before saving"), 409
    if not _advance_collection_revision(state):
        return jsonify(error="Collection changed during save; recovery is required before further edits"), 503
    activity(g.user, "delete", "collection", collection_id)
    return jsonify(ok=True)


@api.get("/linesheets/dashboard")
@auth_required
def linesheet_dashboard():
    totals = {"total": db().linesheets.count_documents({})}
    for status in ("draft", "confirmed", "converted_to_order", "archived"):
        totals[status] = db().linesheets.count_documents({"status": status})
    by_collection = list(db().linesheets.aggregate([{"$group": {"_id": "$collection", "count": {"$sum": 1}}}, {"$sort": {"count": -1}}]))
    by_client = list(db().linesheets.aggregate([{"$group": {"_id": "$client_name", "count": {"$sum": 1}}}, {"$sort": {"count": -1}}, {"$limit": 10}]))
    sync = list(db().linesheets.aggregate([{"$group": {"_id": {"$ifNull": ["$workdrive_status", "not_uploaded"]}, "count": {"$sum": 1}}}]))
    imports = list(db().imports.find({}, {"filename": 1, "status": 1, "created_at": 1, "summary": 1}).sort("created_at", DESCENDING).limit(8))
    return jsonify(serialize({"totals": totals, "by_collection": by_collection, "by_client": by_client, "workdrive": sync, "recent_imports": imports}))


def _linesheet_items_for_consumers(items):
    """Expose legacy/imported fields through the normal consumer item contract."""
    normalized = []
    for item in items or []:
        compatibility = {}
        if "vendor_code" not in item and item.get("source_vendor_code") is not None:
            compatibility["vendor_code"] = item["source_vendor_code"]
        if "size_quantities" not in item and item.get("source_order_quantity") is not None:
            compatibility["size_quantities"] = item["source_order_quantity"]
        if not compatibility:
            normalized.append(item)
            continue
        normalized.append({**item, **compatibility})
    return normalized


@api.route("/linesheets/<linesheet_id>", methods=["GET", "PATCH"])
@auth_required
def linesheet_detail(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    if request.method == "PATCH":
        if g.user["role"] not in {"admin", "sales"}:
            return jsonify(error="Permission denied"), 403
        changes = body()
        if "items" in changes and sheet.get("status") != "draft":
            return jsonify(error="Only draft linesheets can be edited"), 409
        allowed = {k: v for k, v in changes.items() if k in {"name", "collection", "status", "items"}}
        collection = _collection_for_name(allowed.get("collection") or sheet.get("collection"))
        if "items" in allowed:
            allowed["items"] = _normalize_linesheet_items(allowed["items"], collection)
            allowed.update(total_quantity=0, total_value=money(0))
        allowed.update(collection=collection["name"], collection_id=collection["_id"], collection_code=collection["code"])
        allowed["updated_at"] = now()
        db().linesheets.update_one({"_id": sheet["_id"]}, {"$set": allowed})
        activity(g.user, "update", "linesheet", linesheet_id, {"fields": list(allowed)})
        sheet = db().linesheets.find_one({"_id": sheet["_id"]})
        if _is_mds_linesheet(sheet):
            _sync_mds_linesheet_workdrive(sheet)
            sheet = db().linesheets.find_one({"_id": sheet["_id"]})
    related_order = db().orders.find_one({"source_linesheet_id": sheet["_id"]})
    documents = list(db().documents.find({"linesheet_id": sheet["_id"]}).sort("created_at", DESCENDING))
    sheet = {**sheet, "items": _linesheet_items_for_consumers(sheet.get("items", []))}
    return jsonify(serialize({"linesheet": sheet, "related_order": related_order, "documents": documents}))


@api.post("/linesheets/<linesheet_id>/images")
@permission_required("linesheets:write")
def upload_linesheet_image(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    uploaded = request.files.get("file")
    product_code = (request.form.get("product_code") or "").strip()
    if not sheet or not uploaded or not product_code:
        return jsonify(error="Linesheet, product code and image are required"), 400
    allowed = {"image/jpeg", "image/png", "image/webp"}
    if uploaded.mimetype not in allowed:
        return jsonify(error="JPG, PNG or WebP images are required"), 400
    data = uploaded.read()
    if len(data) > 10 * 1024 * 1024:
        return jsonify(error="Image exceeds the 10 MB limit"), 413
    folder = (Path(current_app.config["UPLOAD_DIR"]) / "linesheets" / str(sheet["_id"])).resolve(); folder.mkdir(parents=True, exist_ok=True)
    filename = f"{uuid4()}-{secure_filename(uploaded.filename or 'image')}"; path = folder / filename; path.write_bytes(data)
    image = {"id": str(uuid4()), "path": str(path), "filename": secure_filename(uploaded.filename or "image"), "content_type": uploaded.mimetype, "created_at": now()}
    result = db().linesheets.update_one({"_id": sheet["_id"], "items.product_code": product_code}, {"$set": {"items.$.images": [image], "updated_at": now()}})
    if not result.modified_count:
        return jsonify(error="Product code was not found on this linesheet"), 404
    return jsonify(image={key: value for key, value in image.items() if key != "path"}), 201


@api.post("/linesheets/<linesheet_id>/duplicate")
@permission_required("linesheets:write")
def duplicate_linesheet(linesheet_id):
    source = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not source:
        return jsonify(error="Linesheet not found"), 404
    duplicate = {key: value for key, value in source.items() if key not in {"_id", "order_id", "client_id", "client_name", "created_at", "updated_at"}}
    duplicate.update(name=f"{source.get('name') or source['linesheet_number']} Copy", linesheet_number=next_number(db(), "linesheet", "LS"), status="draft", created_by=g.user["_id"], created_at=now(), updated_at=now())
    result = db().linesheets.insert_one(duplicate)
    activity(g.user, "duplicate", "linesheet", result.inserted_id, {"source_id": linesheet_id})
    return jsonify(id=str(result.inserted_id), linesheet_number=duplicate["linesheet_number"]), 201


@api.post("/linesheets/<linesheet_id>/share")
@permission_required("linesheets:write")
def share_linesheet(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    try:
        data = body(("recipient",))
        share_record = {"recipient": str(data["recipient"]).strip(), "shared_by": g.user["_id"], "shared_at": now()}
        db().linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"status": "shared", "last_shared": share_record, "updated_at": now()}, "$push": {"share_history": share_record}})
        activity(g.user, "share", "linesheet", linesheet_id, {"recipient": share_record["recipient"]})
        return jsonify(ok=True, status="shared")
    except ValueError as exc:
        return json_error(exc)


@api.route("/linesheets/<linesheet_id>/purchase-orders", methods=["GET", "POST"])
@auth_required
def linesheet_purchase_orders(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    if not _is_mds_linesheet(sheet):
        return jsonify(error="Automatic PO folders are available only for MDS linesheets", eligible=False), 409
    try:
        folder = _ensure_mds_po_folder(sheet)
    except RuntimeError as exc:
        return jsonify(error=str(exc), retryable=True), 503
    order = db().orders.find_one({"source_linesheet_id": sheet["_id"]})
    if request.method == "GET":
        query = {"kind": "purchase_order", "$or": [{"linesheet_id": sheet["_id"]}] + ([{"order_id": order["_id"]}] if order else [])}
        items = list(db().documents.find(query).sort("created_at", DESCENDING))
        return jsonify(eligible=True, folder=serialize({key: value for key, value in folder.items() if key != "path"}), order=serialize(order) if order else None, items=serialize(items))
    if g.user["role"] not in {"admin", "sales"}:
        return jsonify(error="Permission denied"), 403
    uploads = request.files.getlist("files") or request.files.getlist("file")
    if not uploads:
        return jsonify(error="Select at least one PDF purchase order"), 400
    prepared = []
    for uploaded in uploads:
        data = uploaded.read()
        if uploaded.mimetype != "application/pdf" or not data.startswith(b"%PDF"):
            return jsonify(error=f"{uploaded.filename or 'File'} is not a valid PDF"), 400
        prepared.append((uploaded, data))
    created = []
    for uploaded, data in prepared:
        document_id = ObjectId(); filename = secure_filename(uploaded.filename or "purchase-order.pdf") or "purchase-order.pdf"
        path = Path(folder["path"]) / f"{document_id}-{filename}"; path.write_bytes(data)
        doc = {"_id": document_id, "family_id": str(document_id), "version": 1, "kind": "purchase_order", "linesheet_id": sheet["_id"], "order_id": order["_id"] if order else None, "client_id": sheet.get("client_id"), "po_folder_id": folder["id"], "original_name": filename, "stored_path": str(path), "content_type": "application/pdf", "size": len(data), "sha256": hashlib.sha256(data).hexdigest(), "uploaded_by": g.user["_id"], "uploaded_by_email": g.user["email"], "created_at": now()}
        workdrive = WorkDriveClient()
        if workdrive.configured:
            try:
                _, workdrive_po_folder = _ensure_mds_workdrive_folders(sheet, workdrive)
                item = _workdrive_uploaded_item(workdrive.upload_bytes(data, workdrive_po_folder, filename, "application/pdf"))
                if not item.get("id"):
                    raise WorkDriveError("WorkDrive did not confirm the PO upload", "file_upload_failure")
                doc.update({"upload_status": "synced", "storage_provider": "local_and_workdrive", "workdrive_folder_id": workdrive_po_folder, "workdrive_file_id": item["id"], "workdrive_link": (item.get("attributes") or {}).get("permalink"), "workdrive_uploaded_at": now()})
            except WorkDriveError as exc:
                current_app.logger.warning("PO WorkDrive sync failed: document_id=%s kind=%s", document_id, exc.kind)
                doc.update({"upload_status": "failed", "sync_error": str(exc), "sync_error_kind": exc.kind})
        else:
            doc["upload_status"] = "pending"
        db().documents.insert_one(doc); created.append(doc)
    activity(g.user, "upload_purchase_orders", "linesheet", linesheet_id, {"count": len(created)})
    return jsonify(items=serialize(created)), 201


@api.post("/linesheets/<linesheet_id>/sync-workdrive")
@permission_required("linesheets:write")
def sync_linesheet_workdrive(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    if not _is_mds_linesheet(sheet):
        return jsonify(error="WorkDrive auto-sync applies only to MDS linesheets"), 409
    result = _sync_mds_linesheet_workdrive(sheet)
    if result["status"] == "failed":
        return jsonify(error=result.get("error"), status="failed", retryable=True), 503
    return jsonify(ok=result["status"] == "synced", status=result["status"]), 200 if result["status"] == "synced" else 202


@api.delete("/linesheets/<linesheet_id>")
@permission_required("linesheets:write")
def delete_linesheet(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    if sheet.get("status") != "draft" or db().orders.find_one({"source_linesheet_id": sheet["_id"]}):
        return jsonify(error="Only unlinked draft linesheets can be deleted"), 409
    db().linesheets.delete_one({"_id": sheet["_id"]})
    activity(g.user, "delete", "linesheet", linesheet_id)
    return jsonify(ok=True)


@api.get("/linesheets/<linesheet_id>/images/<image_id>")
@auth_required
def linesheet_image(linesheet_id, image_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    image = next((image for item in (sheet or {}).get("items", []) for image in item.get("images", []) if image.get("id") == image_id), None)
    if not image or not os.path.exists(image.get("path", "")):
        return jsonify(error="Image not found"), 404
    return send_file(image["path"], mimetype=image.get("content_type", "image/jpeg"), as_attachment=False)


@api.get("/linesheets/<linesheet_id>/export.pdf")
@auth_required
def export_linesheet_pdf(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    output = BytesIO(_build_linesheet_pdf(sheet))
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name=f"{sheet['linesheet_number']}.pdf")


def _build_linesheet_pdf(sheet):
    sheet = {**sheet, "items": _linesheet_items_for_consumers(sheet.get("items", []))}
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=landscape(A4))
    width, height = landscape(A4)
    pdf.setTitle(sheet.get("name") or sheet["linesheet_number"])
    pdf.setFont("Helvetica-Bold", 20); pdf.drawString(36, height - 42, "RK FASHION")
    pdf.setFont("Helvetica-Bold", 14); pdf.drawString(36, height - 68, sheet.get("name") or sheet["linesheet_number"])
    pdf.setFont("Helvetica", 10); pdf.drawString(36, height - 85, f"{sheet.get('collection', '')} | {sheet['linesheet_number']} | {sheet.get('status', 'draft').upper()}")
    y = height - 120
    headers = ["Image", "Product code", "Description", "Color", "Set", "Sizes", "MRP", "SKU"]
    positions = [36, 92, 175, 310, 382, 420, 510, 565]
    pdf.setFont("Helvetica-Bold", 9)
    for label, x in zip(headers, positions): pdf.drawString(x, y, label)
    y -= 18; pdf.setFont("Helvetica", 8)
    for item in sheet.get("items", []):
        if y < 40:
            pdf.showPage(); y = height - 42
        image_paths = item.get("images") or []
        image_path = image_paths[0].get("path") if image_paths and isinstance(image_paths[0], dict) else (image_paths[0] if image_paths else None)
        if image_path and os.path.exists(image_path):
            try: pdf.drawImage(ImageReader(image_path), positions[0], y - 28, width=42, height=42, preserveAspectRatio=True, anchor="c")
            except Exception: pass
        values = [item.get("product_code") or item.get("vendor_code"), item.get("description") or item.get("product_name"), item.get("color"), item.get("set_of", 1), ", ".join(item.get("sizes") or item.get("size_quantities", {}).keys()), str(item.get("mrp") or item.get("unit_price") or 0), item.get("linesheet_sku") or item.get("sku")]
        for value, x in zip(values, positions[1:]): pdf.drawString(x, y, str(value or "")[:28])
        y -= 48
    pdf.save()
    return output.getvalue()


@api.post("/linesheets/<linesheet_id>/email")
@permission_required("linesheets:write")
def email_linesheet(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    try:
        data = body(("to", "subject", "body", "attachments", "request_id"))
    except ValueError as exc:
        return json_error(exc)
    recipient = str(data["to"]).strip()
    email_pattern = re.compile(r"^[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+$")
    def email_list(value):
        return [item.strip() for item in re.split(r"[,;]", str(value or "")) if item.strip()]
    recipients, cc_recipients, bcc_recipients = email_list(recipient), email_list(data.get("cc")), email_list(data.get("bcc"))
    invalid = [item for item in recipients + cc_recipients + bcc_recipients if not email_pattern.match(item)]
    if len(recipients) != 1 or invalid:
        return jsonify(error=f"Invalid email address: {invalid[0] if invalid else recipient}"), 400
    attachment_kind = str(data["attachments"]).lower()
    if attachment_kind not in {"excel", "pdf", "both"}:
        return jsonify(error="Attachments must be excel, pdf or both"), 400
    request_id = str(data["request_id"]).strip()
    existing = db().linesheet_email_history.find_one({"request_id": request_id})
    if existing:
        return jsonify(serialize({"ok": True, "duplicate": True, "history": existing}))
    cc_value, bcc_value = ",".join(cc_recipients), ",".join(bcc_recipients)
    history = {"request_id": request_id, "linesheet_id": sheet["_id"], "to": recipient, "cc": cc_value, "bcc": bcc_value, "attachments": attachment_kind, "status": "sending", "created_at": now(), "created_by": g.user["_id"]}
    try:
        db().linesheet_email_history.insert_one(history)
    except DuplicateKeyError:
        existing = db().linesheet_email_history.find_one({"request_id": request_id})
        return jsonify(serialize({"ok": True, "duplicate": True, "history": existing}))
    attachments = []
    try:
        if attachment_kind in {"pdf", "both"}:
            attachments.append((_linesheet_attachment_name(sheet, "pdf"), _build_linesheet_pdf(sheet), "application/pdf"))
        if attachment_kind in {"excel", "both"}:
            attachments.append((_linesheet_attachment_name(sheet, "xlsx"), _build_linesheet_excel(sheet), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        result = ZohoMailClient().send_message(recipient, str(data["subject"]).strip(), str(data["body"]), attachments=attachments, cc=cc_value, bcc=bcc_value)
        db().linesheet_email_history.update_one({"request_id": request_id}, {"$set": {"status": "sent", "provider_message_id": result.get("provider_message_id"), "sent_at": now()}})
        return jsonify(ok=True, status="sent")
    except (WorkDriveError, ValueError) as exc:
        kind = exc.kind if isinstance(exc, WorkDriveError) else "validation_error"
        current_app.logger.warning("Linesheet email failed: kind=%s linesheet_id=%s", kind, linesheet_id)
        db().linesheet_email_history.update_one({"request_id": request_id}, {"$set": {"status": "failed", "error": str(exc), "failed_at": now()}})
        return jsonify(error=str(exc), status=kind), 503 if isinstance(exc, WorkDriveError) else 400


def _linesheet_attachment_name(sheet, extension):
    stem = secure_filename(f"{sheet.get('linesheet_number', 'linesheet')}-{sheet.get('name', '')}").strip("-_") or "linesheet"
    return f"{stem}.{extension}"


@api.get("/linesheets/<linesheet_id>/email-history")
@permission_required("linesheets:read")
def linesheet_email_history(linesheet_id):
    rows = db().linesheet_email_history.find({"linesheet_id": oid(linesheet_id)}, {"request_id": 0, "created_by": 0, "error": 0}).sort("created_at", DESCENDING)
    return jsonify(items=serialize(list(rows)))


def _normalized_header(value):
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _excel_mapping(headers):
    aliases = {"serial_number": {"sr no", "sr. no", "serial number"}, "vendor_code": {"vendor code"}, "sku": {"sku", "product sku"}, "color": {"color", "colour"}, "category": {"category"}, "component_count": {"no of components", "component count"}, "mrp": {"mrp"}, "remark": {"remark", "remarks"}, "reference_image_1": {"reference image 1"}, "reference_image_2": {"reference image 2"}, "po_delivery_date": {"po deliverydate", "po delivery date", "delivery date"}}
    mapping, sizes = {}, []
    for index, header in enumerate(headers):
        normalized = _normalized_header(header)
        for field, names in aliases.items():
            if normalized in names:
                mapping[field] = index
        first = normalized.split(" ")[0].upper() if normalized else ""
        if first in {"XXS", "XS", "S", "M", "L", "XL", "XXL", "XXXL", "FREE", "OS"}:
            measurements = {}
            for label, number in re.findall(r"([A-Za-z][A-Za-z ]*)\s*:\s*(\d+(?:\.\d+)?)", str(header or "")):
                measurements[label.strip().lower()] = float(number) if "." in number else int(number)
            sizes.append({"name": first, "column": index, "header": str(header), "measurements": measurements})
    return mapping, sizes


def _parse_quantity(value):
    """Parse the leading quantity while retaining any order reference text."""
    if value is None or value == "":
        return 0, ""
    if isinstance(value, (int, float)):
        quantity = float(value)
        return (int(quantity) if quantity.is_integer() else quantity), ""
    text = str(value).strip()
    match = re.match(r"^\s*(-?\d+(?:\.\d+)?)", text)
    if not match:
        raise ValueError("quantity is not numeric")
    quantity = float(match.group(1))
    reference = " ".join(line.strip() for line in text[match.end():].splitlines() if line.strip()).strip(" ()")
    return (int(quantity) if quantity.is_integer() else quantity), reference


@api.post("/linesheets/import/inspect")
@permission_required("linesheets:write")
def inspect_linesheet_import():
    uploaded = request.files.get("file")
    if not uploaded or not (uploaded.filename or "").lower().endswith(".xlsx"):
        return jsonify(error="An .xlsx workbook is required"), 400
    raw = uploaded.read()
    if len(raw) > 25 * 1024 * 1024:
        return jsonify(error="Workbook exceeds the 25 MB limit"), 413
    import_id = str(uuid4())
    folder = (Path(current_app.config["UPLOAD_DIR"]) / "imports" / import_id).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    source_path = folder / secure_filename(uploaded.filename or "linesheet.xlsx")
    source_path.write_bytes(raw)
    file_hash = hashlib.sha256(raw).hexdigest()
    try:
        workbook = load_workbook(source_path, data_only=True)
    except Exception:
        return jsonify(error="Workbook could not be opened"), 400
    worksheet_name = request.form.get("worksheet") or workbook.sheetnames[0]
    if worksheet_name not in workbook.sheetnames:
        return jsonify(error="Worksheet not found", worksheets=workbook.sheetnames), 400
    worksheet = workbook[worksheet_name]
    header_row = 1
    best_headers = []
    best_score = -1
    for candidate in range(1, min(worksheet.max_row, 25) + 1):
        candidate_headers = [cell.value for cell in worksheet[candidate]]
        candidate_mapping, candidate_sizes = _excel_mapping(candidate_headers)
        score = len(candidate_mapping) + len(candidate_sizes)
        if score > best_score:
            best_score, best_headers, header_row = score, candidate_headers, candidate
    headers = best_headers
    mapping, size_columns = _excel_mapping(headers)
    preview, errors, seen, record_references = [], [], set(), {}
    embedded_by_row = {}
    for number, image in enumerate(getattr(worksheet, "_images", []), 1):
        try:
            row_number = image.anchor._from.row + 1
            image_path = folder / f"embedded-{number}.png"
            image_path.write_bytes(image._data())
            embedded_by_row.setdefault(row_number, []).append(str(image_path))
        except Exception:
            continue
    for row_number, cells in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        if not any(value not in (None, "") for value in cells):
            continue
        sku = str(cells[mapping["sku"]]).strip() if mapping.get("sku") is not None and cells[mapping["sku"]] is not None else ""
        row_errors = []
        if not sku:
            row_errors.append("Missing SKU")
        if sku and sku in seen:
            row_errors.append("Duplicate SKU")
        seen.add(sku)
        quantities, row_references = {}, []
        for size in size_columns:
            value = cells[size["column"]]
            try:
                quantities[size["name"]], reference = _parse_quantity(value)
                if quantities[size["name"]] < 0:
                    raise ValueError()
                if reference:
                    row_references.append(reference)
            except (ValueError, TypeError):
                row_errors.append(f"Invalid quantity for {size['name']}")
        record = {"row_number": row_number, "values": [serialize(v) for v in cells], "sku": sku, "size_quantities": quantities, "size_measurements": {size["name"]: size.get("measurements", {}) for size in size_columns}, "references": row_references, "images": embedded_by_row.get(row_number, []), "errors": row_errors}
        preview.append(record)
        errors.extend({"row": row_number, "message": message} for message in row_errors)
    doc = {"import_id": import_id, "filename": source_path.name, "file_hash": file_hash, "source_path": str(source_path), "worksheet": worksheet_name, "header_row": header_row, "worksheets": workbook.sheetnames, "headers": [str(v or "") for v in headers], "mapping": mapping, "size_columns": size_columns, "preview": preview, "status": "preview", "created_by": g.user["_id"], "created_at": now()}
    db().imports.insert_one(doc)
    return jsonify(serialize({"import_id": import_id, "filename": source_path.name, "worksheets": workbook.sheetnames, "worksheet": worksheet_name, "header_row": header_row, "headers": doc["headers"], "mapping": mapping, "size_columns": size_columns, "rows": preview[:100], "summary": {"total": len(preview), "valid": len(preview)-len({e['row'] for e in errors}), "invalid": len({e['row'] for e in errors}), "errors": len(errors)}}))


@api.get("/linesheets/import/<import_id>/aakaar-dry-run")
@permission_required("linesheets:write")
def aakaar_import_dry_run(import_id):
    imported = db().imports.find_one({"import_id": import_id})
    if not imported:
        return jsonify(error="Import preview not found"), 404
    collection = db().collections.find_one({"name": "Aakaar", "slug": "aakaar", "code": "AAK", "active": True})
    projection = build_aakaar_projection(imported, collection, existing_aakaar_identities(db()))
    return jsonify(serialize({"import_id": import_id, "projection": projection}))


@api.post("/linesheets/import/<import_id>/commit")
@permission_required("linesheets:write")
def commit_linesheet_import(import_id):
    imported = db().imports.find_one({"import_id": import_id})
    if not imported:
        return jsonify(error="Import preview not found"), 404
    try:
        data = body(("client_id", "collection", "type", "name"))
        existing_import = db().imports.find_one({"client_id": oid(data["client_id"]), "file_hash": imported.get("file_hash"), "status": "imported"})
        if existing_import:
            return jsonify(error="This workbook has already been imported for this client", duplicate=True, existing_import_id=existing_import["import_id"], existing_linesheet_id=str(existing_import.get("linesheet_id"))), 409
        if imported.get("status") == "imported" and imported.get("linesheet_id"):
            return jsonify(id=str(imported["linesheet_id"]), duplicate=True, existing_import_id=import_id), 200
        client = db().clients.find_one({"_id": oid(data["client_id"])})
        collection = db().collections.find_one({"name": data["collection"], "active": True})
        if not client or not collection:
            raise ValueError("Valid client and collection are required")
        if data.get("model") == "aakaar":
            projection = build_aakaar_projection(imported, collection, existing_aakaar_identities(db()))
            if not projection["valid"]:
                return jsonify(error="Aakaar validation failed; no records were written", projection=serialize(projection)), 400
            result = commit_aakaar_projection(db(), projection, imported, collection, client, data, g.user)
            db().imports.update_one({"_id": imported["_id"]}, {"$set": {"status": "imported", "client_id": client["_id"], "linesheet_id": result["linesheet_id"], "summary": result, "completed_at": now()}})
            activity(g.user, "import", "aakaar", result["linesheet_id"], result)
            return jsonify(id=str(result["linesheet_id"]), summary=serialize(result)), 201
        mapping, items, skipped, created_products = imported["mapping"], [], 0, 0
        for row in imported["preview"]:
            if row["errors"]:
                skipped += 1
                continue
            values = row["values"]
            def value(field, default=None):
                index = mapping.get(field)
                return values[index] if index is not None and index < len(values) else default
            sku = str(row["sku"])
            product = db().products.find_one({"sku": sku})
            item = {"sku": sku, "vendor_code": str(value("vendor_code", "") or ""), "product_name": product.get("name") if product else sku, "color": value("color", ""), "category": value("category", ""), "component_count": value("component_count", 0), "mrp": money(value("mrp", 0)), "unit_price": money(value("mrp", 0)), "remark": " ".join(filter(None, [str(value("remark", "") or ""), *row.get("references", [])])), "po_reference": (row.get("references") or [None])[0], "po_delivery_date": value("po_delivery_date"), "size_measurements": row.get("size_measurements", {}), "size_quantities": row["size_quantities"], "images": row["images"]}
            items.append(item)
            if not product and data.get("create_products", False):
                db().products.insert_one({"name": item["product_name"], "sku": sku, "vendor_code": item["vendor_code"], "collection": data["collection"], "category": item["category"], "colors": [item["color"]], "sizes": list(item["size_quantities"]), "selling_price": item["mrp"], "images": row["images"], "active": True, "created_at": now(), "updated_at": now()})
                created_products += 1
        total_qty = sum(sum(int(v) for v in item["size_quantities"].values()) for item in items)
        total_value = sum(sum(int(v) for v in item["size_quantities"].values()) * float(item["unit_price"]) for item in items)
        normalized_title = re.sub(r"\s+", " ", data["name"]).strip().casefold()
        if data.get("reference_number") and db().linesheets.find_one({"client_id": client["_id"], "collection": data["collection"], "type": data["type"], "reference_number": data["reference_number"]}):
            return jsonify(error="A linesheet with this client, collection, order type and reference already exists", duplicate=True), 409
        if not data.get("reference_number") and db().linesheets.find_one({"client_id": client["_id"], "name_normalized": normalized_title, "source_file_hash": imported.get("file_hash")}):
            return jsonify(error="A linesheet with this title and workbook already exists for this client", duplicate=True), 409
        sheet = {"linesheet_number": next_number(db(), "linesheet", "LS"), "name": data["name"], "name_normalized": normalized_title, "reference_number": data.get("reference_number"), "source_file_hash": imported.get("file_hash"), "collection": data["collection"], "client_id": client["_id"], "client_name": client["name"], "type": data["type"], "items": items, "total_quantity": total_qty, "total_value": money(total_value), "status": data.get("status", "draft"), "source_import_id": import_id, "original_workbook": {"filename": imported["filename"], "path": imported["source_path"], "worksheet": imported["worksheet"], "headers": imported["headers"]}, "workdrive_status": "not_uploaded", "created_by": g.user["_id"], "created_at": now(), "updated_at": now()}
        result = db().linesheets.insert_one(sheet)
        sheet["_id"] = result.inserted_id
        try:
            _ensure_mds_po_folder(sheet)
        except RuntimeError:
            db().linesheets.delete_one({"_id": result.inserted_id})
            raise
        backup = {"linesheet_id": result.inserted_id, "linesheet_title": sheet["name"], "client_id": client["_id"], "client_name": client["name"], "original_filename": imported["filename"], "source_path": imported["source_path"], "import_id": import_id, "uploaded_by": g.user["_id"], "uploaded_at": now(), "status": "pending"}
        sync = _sync_mds_linesheet_workdrive(sheet)
        original_sync = (sync.get("files") or {}).get("original") or {}
        backup.update({"status": "uploaded" if original_sync.get("file_id") else sync["status"], "workdrive_folder_id": (db().linesheets.find_one({"_id": result.inserted_id}) or {}).get("workdrive_folder_id"), "workdrive_file_id": original_sync.get("file_id"), "error": sync.get("error")})
        db().workdrive_original_uploads.insert_one(backup)
        summary = {"products_imported": len(items), "rows_skipped": skipped, "errors": sum(len(row["errors"]) for row in imported["preview"]), "new_products_created": created_products, "existing_products_matched": len(items)-created_products}
        db().imports.update_one({"_id": imported["_id"]}, {"$set": {"status": "imported", "client_id": client["_id"], "linesheet_id": result.inserted_id, "summary": summary, "completed_at": now()}})
        activity(g.user, "import", "linesheet", result.inserted_id, summary)
        return jsonify(id=str(result.inserted_id), linesheet_number=sheet["linesheet_number"], summary=summary), 201
    except Exception as exc:
        return json_error(exc)


@api.route("/linesheets/<linesheet_id>/original-upload", methods=["GET", "POST"])
@auth_required
def original_workbook_upload(linesheet_id):
    record = db().workdrive_original_uploads.find_one({"linesheet_id": oid(linesheet_id)}, sort=[("uploaded_at", -1)])
    if request.method == "GET":
        return jsonify(item=serialize(record) if record else None)
    if not record:
        return jsonify(error="Original workbook record not found"), 404
    try:
        workdrive = WorkDriveClient()
        folder_id, _ = workdrive.ensure_folder_path(["Client Linesheets", record["client_name"], record["linesheet_title"], "Original Uploads"])
        uploaded_file = workdrive.upload(record["source_path"], folder_id, record["original_filename"])
        data = uploaded_file.get("data", uploaded_file) if isinstance(uploaded_file, dict) else {}
        db().workdrive_original_uploads.update_one({"_id": record["_id"]}, {"$set": {"status": "uploaded", "workdrive_folder_id": folder_id, "workdrive_file_id": data.get("id"), "retry_at": now()}, "$unset": {"error": "", "error_kind": ""}})
        return jsonify(ok=True, status="uploaded", workdrive_file_id=data.get("id"))
    except WorkDriveError as exc:
        db().workdrive_original_uploads.update_one({"_id": record["_id"]}, {"$set": {"status": "failed", "error": str(exc), "error_kind": exc.kind, "retry_at": now()}})
        return jsonify(error=str(exc), status=exc.kind), 503


@api.get("/linesheets/<linesheet_id>/export.xlsx")
@auth_required
def export_linesheet_excel(linesheet_id):
    sheet_doc = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet_doc:
        return jsonify(error="Linesheet not found"), 404
    original = request.args.get("format") == "original" and sheet_doc.get("original_workbook")
    output = BytesIO(_build_linesheet_excel(sheet_doc, original=original))
    db().documents.insert_one({"family_id": str(uuid4()), "version": 1, "linesheet_id": sheet_doc["_id"], "kind": "linesheet_excel_export", "original_name": f"{sheet_doc['linesheet_number']}.xlsx", "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "size": output.getbuffer().nbytes, "storage_provider": "generated", "upload_status": "not_uploaded", "uploaded_by": g.user["_id"], "created_at": now()})
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"{sheet_doc['linesheet_number']}.xlsx")


def _build_linesheet_excel(sheet_doc, original=False):
    sheet_doc = {**sheet_doc, "items": _linesheet_items_for_consumers(sheet_doc.get("items", []))}
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = (sheet_doc.get("collection") or "Linesheet")[:31]
    size_names = []
    for item in sheet_doc.get("items", []):
        for size in item.get("size_quantities", {}):
            if size not in size_names:
                size_names.append(size)
    standard_headers = ["Sr No", "Image", "Vendor Code", "Sku", "Color", "Category", "No of Components", "MRP", *size_names, "Remark", "Reference Image 1", "Reference Image 2", "Po DeliveryDate"]
    headers = sheet_doc.get("original_workbook", {}).get("headers") if original else standard_headers
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="2A2622")
    image_bytes_cache = {}
    for index, item in enumerate(sheet_doc.get("items", []), 1):
        values = {"Sr No": index, "Vendor Code": item.get("vendor_code"), "Sku": item.get("sku"), "Color": item.get("color"), "Category": item.get("category"), "No of Components": item.get("component_count"), "MRP": float(item.get("mrp") or item.get("unit_price") or 0), "Remark": item.get("remark"), "Reference Image 1": item.get("reference_image_1"), "Reference Image 2": item.get("reference_image_2"), "Po DeliveryDate": item.get("po_delivery_date")}
        values.update(item.get("size_quantities", {}))
        worksheet.append([values.get(header, values.get(str(header).split(" ")[0].upper(), "")) for header in headers])
        images = item.get("images") or []
        image_path = images[0].get("path") if images and isinstance(images[0], dict) else (images[0] if images else None)
        if image_path and os.path.exists(image_path) and "Image" in headers:
            try:
                if image_path not in image_bytes_cache:
                    image_bytes_cache[image_path] = Path(image_path).read_bytes()
                excel_image = ExcelImage(BytesIO(image_bytes_cache[image_path]))
                ratio = min(72 / max(excel_image.width, 1), 88 / max(excel_image.height, 1))
                excel_image.width = max(1, round(excel_image.width * ratio)); excel_image.height = max(1, round(excel_image.height * ratio))
                image_column = headers.index("Image") + 1
                worksheet.add_image(excel_image, f"{get_column_letter(image_column)}{index+1}")
                worksheet.row_dimensions[index+1].height = max(70, excel_image.height * 0.75)
                worksheet.column_dimensions[get_column_letter(image_column)].width = max(14, worksheet.column_dimensions[get_column_letter(image_column)].width or 0)
            except (OSError, TypeError, ValueError):
                current_app.logger.warning("Skipped invalid linesheet image during Excel export: linesheet_id=%s row=%s", sheet_doc.get("_id"), index)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in worksheet.columns:
        letter = column[0].column_letter
        calculated = min(max(12, max(len(str(cell.value or "")) for cell in column) + 2), 35)
        worksheet.column_dimensions[letter].width = max(worksheet.column_dimensions[letter].width or 0, calculated)
    output = BytesIO(); workbook.save(output)
    return output.getvalue()


@api.get("/linesheets/<linesheet_id>/preview.xlsx")
@auth_required
def preview_linesheet_excel(linesheet_id):
    sheet = db().linesheets.find_one({"_id": oid(linesheet_id)})
    if not sheet:
        return jsonify(error="Linesheet not found"), 404
    workbook = load_workbook(BytesIO(_build_linesheet_excel(sheet)), data_only=True)
    worksheets = []
    for worksheet in workbook.worksheets:
        rows = [[serialize(cell.value) for cell in row] for row in worksheet.iter_rows()]
        widths = {key: value.width for key, value in worksheet.column_dimensions.items() if value.width}
        worksheets.append({"name": worksheet.title, "rows": rows, "freeze_panes": str(worksheet.freeze_panes or ""), "column_widths": widths})
    return jsonify(linesheet=serialize({"_id": sheet["_id"], "name": sheet.get("name"), "linesheet_number": sheet.get("linesheet_number")}), worksheets=worksheets)


def _workdrive_uploaded_item(payload):
    data = payload.get("data", payload) if isinstance(payload, dict) else {}
    if isinstance(data, list):
        data = data[0] if data else {}
    return data if isinstance(data, dict) else {}


def _ensure_mds_workdrive_folders(sheet, client=None):
    client = client or WorkDriveClient()
    mds_id = client.ensure_managed_folder("mds-root", "MDS", client.root_folder_id)
    label = f"{sheet.get('linesheet_number', '')} - {sheet.get('name') or sheet.get('title') or sheet['_id']}".strip(" -")
    sheet_id = client.ensure_managed_folder(f"mds-linesheet:{sheet['_id']}", label, mds_id)
    po_id = client.ensure_managed_folder(f"mds-linesheet:{sheet['_id']}:purchase-orders", "Purchase Orders", sheet_id)
    db().linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"workdrive_folder_id": sheet_id, "workdrive_po_folder_id": po_id, "workdrive_folder_name": label}})
    return sheet_id, po_id


def _sync_mds_linesheet_workdrive(sheet):
    if not _is_mds_linesheet(sheet):
        return {"status": "not_applicable"}
    client = WorkDriveClient()
    if not client.configured:
        db().linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"workdrive_status": "pending", "workdrive_error": "WorkDrive is not configured", "updated_at": now()}})
        return {"status": "pending"}
    try:
        folder_id, po_folder_id = _ensure_mds_workdrive_folders(sheet, client)
        stem = secure_filename(f"{sheet.get('linesheet_number', 'linesheet')}-{sheet.get('name', '')}").strip("-_") or "linesheet"
        uploads = {}
        generated = [("excel", f"{stem}.xlsx", _build_linesheet_excel(sheet), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), ("pdf", f"{stem}.pdf", _build_linesheet_pdf(sheet), "application/pdf")]
        original = sheet.get("original_workbook") or {}
        if original.get("path") and os.path.exists(original["path"]):
            generated.insert(0, ("original", original.get("filename") or "original.xlsx", Path(original["path"]).read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        for kind, filename, content, content_type in generated:
            item = _workdrive_uploaded_item(client.upload_bytes(content, folder_id, filename, content_type))
            if not item.get("id"):
                raise WorkDriveError(f"WorkDrive did not confirm the {kind} upload", "file_upload_failure")
            uploads[kind] = {"file_id": item["id"], "filename": filename, "uploaded_at": now(), "status": "synced", "permalink": (item.get("attributes") or {}).get("permalink")}
        db().linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"workdrive_status": "synced", "workdrive_folder_id": folder_id, "workdrive_po_folder_id": po_folder_id, "workdrive_files": uploads, "workdrive_synced_at": now(), "updated_at": now()}, "$unset": {"workdrive_error": ""}})
        return {"status": "synced", "files": uploads}
    except WorkDriveError as exc:
        current_app.logger.warning("MDS WorkDrive sync failed: linesheet_id=%s kind=%s", sheet["_id"], exc.kind)
        db().linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"workdrive_status": "failed", "workdrive_error": str(exc), "workdrive_error_kind": exc.kind, "workdrive_failed_at": now(), "updated_at": now()}})
        return {"status": "failed", "error": str(exc), "kind": exc.kind}


def _sync_uploaded_mds_linesheet_workdrive(sheet, content):
    client = WorkDriveClient()
    if not client.configured:
        db().client_linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"workdrive_status": "pending"}})
        return {"status": "pending"}
    try:
        mds_id = client.ensure_managed_folder("mds-root", "MDS", client.root_folder_id)
        label = sheet.get("title") or str(sheet["_id"])
        folder_id = client.ensure_managed_folder(f"client-mds-linesheet:{sheet['_id']}", label, mds_id)
        po_folder_id = client.ensure_managed_folder(f"client-mds-linesheet:{sheet['_id']}:purchase-orders", "Purchase Orders", folder_id)
        item = _workdrive_uploaded_item(client.upload_bytes(content, folder_id, sheet["original_filename"], sheet.get("content_type") or "application/octet-stream"))
        if not item.get("id"):
            raise WorkDriveError("WorkDrive did not confirm the original linesheet upload", "file_upload_failure")
        db().client_linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"workdrive_status": "synced", "workdrive_folder_id": folder_id, "workdrive_po_folder_id": po_folder_id, "workdrive_original_file_id": item["id"], "workdrive_synced_at": now()}})
        return {"status": "synced"}
    except WorkDriveError as exc:
        db().client_linesheets.update_one({"_id": sheet["_id"]}, {"$set": {"workdrive_status": "failed", "workdrive_error": str(exc), "workdrive_error_kind": exc.kind}})
        return {"status": "failed", "error": str(exc)}


@api.post("/products/<product_id>/images")
@permission_required("linesheets:write")
def upload_product_images(product_id):
    product = db().products.find_one({"_id": oid(product_id)})
    if not product:
        return jsonify(error="Product not found"), 404
    files = request.files.getlist("files")
    allowed = {"image/jpeg", "image/png", "image/webp"}
    if not files or any(item.mimetype not in allowed for item in files):
        return jsonify(error="JPG, PNG or WebP images are required"), 400
    existing = product.get("images") or []
    asset_folder = product_asset_folder(current_app, product)
    metadata, failures = [], []
    for uploaded in files:
        data = uploaded.read()
        if len(data) > 10 * 1024 * 1024:
            failures.append({"filename": uploaded.filename, "error": "File exceeds 10 MB"}); continue
        if not data or (uploaded.mimetype not in allowed):
            failures.append({"filename": uploaded.filename, "error": "Invalid image file"}); continue
        try:
            remote = upload_image(current_app, BytesIO(data), secure_filename(uploaded.filename or "image"), uploaded.mimetype, asset_folder)
        except (CloudinaryConfigurationError, CloudinaryUploadError) as exc:
            failures.append({"filename": uploaded.filename, "error": str(exc)}); continue
        position = len(existing) + len(metadata)
        metadata.append({"id": str(uuid4()), "url": remote["secure_url"], "secure_url": remote["secure_url"], "public_id": remote.get("public_id"), "asset_folder": asset_folder, "filename": secure_filename(uploaded.filename or "image"), "content_type": uploaded.mimetype, "view": request.form.get("view", "detail"), "description": request.form.get("description", ""), "position": position, "is_main": position == 0 and not existing, "is_primary": position == 0 and not existing, "source": "rk-stock", "created_at": now()})
    if metadata:
        db().products.update_one({"_id": product["_id"]}, {"$push": {"images": {"$each": metadata}}, "$set": {"updated_at": now()}})
        activity(g.user, "upload_images", "product", product_id, {"count": len(metadata), "source": "rk-stock"})
    status = 201 if metadata and not failures else (207 if metadata else 502)
    return jsonify(items=serialize(metadata), failures=failures), status


@api.patch("/products/<product_id>/images")
@permission_required("linesheets:write")
def update_product_images(product_id):
    product = db().products.find_one({"_id": oid(product_id)})
    if not product:
        return jsonify(error="Product not found"), 404
    images = product.get("images") or []
    data = request.get_json(silent=True) or {}
    order = data.get("order")
    primary = data.get("primary_id")
    if order is not None:
        by_id = {str(item.get("id")): item for item in images}
        if not isinstance(order, list) or set(map(str, order)) != set(by_id):
            return jsonify(error="Image order must contain every image exactly once"), 400
        images = [by_id[str(item_id)] for item_id in order]
    if primary is not None and not any(str(item.get("id")) == str(primary) for item in images):
        return jsonify(error="Primary image not found"), 400
    if images:
        primary = str(primary) if primary is not None else next((str(item.get("id")) for item in images if item.get("is_primary") or item.get("is_main")), str(images[0].get("id")))
        for position, item in enumerate(images):
            item["position"] = position; item["is_primary"] = str(item.get("id")) == primary; item["is_main"] = item["is_primary"]
    db().products.update_one({"_id": product["_id"]}, {"$set": {"images": images, "updated_at": now()}})
    return jsonify(items=serialize(images))


@api.delete("/products/<product_id>/images/<image_id>")
@permission_required("linesheets:write")
def remove_product_image(product_id, image_id):
    product = db().products.find_one({"_id": oid(product_id)})
    if not product:
        return jsonify(error="Product not found"), 404
    current_images = product.get("images") or []
    media = next((item for item in current_images if str(item.get("id")) == str(image_id)), None)
    if not media:
        return jsonify(error="Image not found"), 404
    remote_deleted, remote_already_missing = False, False
    if media.get("source") == "rk-stock":
        if not media.get("public_id"):
            return jsonify(error="RK-STOCK media is missing its Cloudinary public ID; no deletion was performed"), 409
        try:
            result = delete_image(current_app, media["public_id"], media.get("resource_type") or "image", media.get("delivery_type") or "upload")
            remote_deleted, remote_already_missing = not result["already_missing"], result["already_missing"]
        except (CloudinaryConfigurationError, CloudinaryUploadError) as exc:
            return jsonify(error=str(exc)), 502
    images = [item for item in current_images if str(item.get("id")) != str(image_id)]
    primary = next((str(item.get("id")) for item in images if item.get("is_primary") or item.get("is_main")), str(images[0].get("id")) if images else None)
    for position, item in enumerate(images):
        item["position"] = position; item["is_primary"] = str(item.get("id")) == primary; item["is_main"] = item["is_primary"]
    db().products.update_one({"_id": product["_id"]}, {"$set": {"images": images, "updated_at": now()}})
    activity(g.user, "remove_product_image", "product", product_id, {"image_id": image_id, "source": media.get("source"), "cloudinary_deleted": remote_deleted, "cloudinary_already_missing": remote_already_missing})
    return jsonify(items=serialize(images), cloudinary_deleted=remote_deleted, cloudinary_already_missing=remote_already_missing, remote_asset_preserved=media.get("source") != "rk-stock")


@api.get("/workdrive/status")
@auth_required
def workdrive_status():
    client = WorkDriveClient()
    stored = db().settings.find_one({"_id": "workdrive_oauth"}) or {}
    verified = bool(stored.get("verified_at") and client.root_folder_configured and stored.get("verified_folder_id") == client.root_folder_id)
    if verified:
        status = "connected"
    elif "token_encryption_key" in client.missing_configuration:
        status = "missing_encryption_key"
    elif "refresh_token" in client.missing_configuration:
        status = "missing_refresh_token"
    else:
        status = stored.get("error_kind") or "incomplete"
    sync_counts = {item["_id"]: item["count"] for item in db().documents.aggregate([{"$group": {"_id": {"$ifNull": ["$upload_status", "not_uploaded"]}, "count": {"$sum": 1}}}])}
    return jsonify(configured=client.configured, missing=client.missing_configuration, root_folder_configured=client.root_folder_configured, verified=verified, status=status, verification_error=stored.get("verification_error"), connected_at=stored.get("connected_at") if verified else None, verified_at=stored.get("verified_at") if verified else None, connected_account_email=stored.get("account_email"), root_folder_name=stored.get("root_folder_name") if verified else None, team_id=stored.get("team_id") if verified else None, storage_structure=([{"name": stored.get("root_folder_name"), "type": "team_folder", "verified": True}] if verified and stored.get("root_folder_name") else []), synchronization={"counts": sync_counts, "last_successful_sync": stored.get("last_successful_sync")})


@api.get("/integrations/storefront/status")
@auth_required
def storefront_integration_status():
    return jsonify(StorefrontIntegrationClient().check())


@api.post("/integrations/storefront/connect")
@permission_required("settings:write")
def storefront_integration_connect():
    try:
        return jsonify(StorefrontIntegrationClient().connect(g.user["_id"]))
    except StorefrontIntegrationError as exc:
        payload = {"error": str(exc), "status": "connection_failed", "error_kind": exc.kind}
        if getattr(exc, "missing", None):
            payload["missing"] = exc.missing
        return jsonify(payload), 503 if exc.kind in {"unreachable", "configuration_error"} else 502 if exc.kind == "credential_rejected" else 409


@api.post("/integrations/storefront/disconnect")
@permission_required("settings:write")
def storefront_integration_disconnect():
    try:
        return jsonify(StorefrontIntegrationClient().disconnect(g.user["_id"]))
    except StorefrontIntegrationError as exc:
        return jsonify(error=str(exc), status="disconnect_failed", error_kind=exc.kind), 503 if exc.kind == "unreachable" else 502 if exc.kind == "credential_rejected" else 409


@api.post("/integrations/storefront/sync-catalog")
@permission_required("settings:write")
def storefront_catalog_sync():
    client = StorefrontIntegrationClient()
    try:
        result = CatalogSyncService(client).sync(g.user["_id"])
        activity(g.user, "sync", "rk-web-catalog", client.record_id, {"products": result["products"], "categories": result["categories"], "collections": result["collections"]})
        return jsonify(serialize(result)), 200
    except StorefrontIntegrationError as exc:
        return jsonify(error=str(exc), status="failed", error_kind=exc.kind), 503 if exc.kind in {"unreachable", "configuration_error"} else 502
    except CatalogSyncError as exc:
        return jsonify(error=str(exc), status="failed", error_kind="malformed_response"), 502
    except Exception:
        current_app.logger.exception("RK-WEB catalog synchronization failed")
        timestamp = now()
        db().settings.update_one({"_id": client.record_id}, {"$set": {"catalog_sync_status": "failed", "catalog_last_error": "Catalog synchronization failed", "updated_at": timestamp}}, upsert=True)
        return jsonify(error="Catalog synchronization failed", status="failed", error_kind="database_failure"), 500


@api.post("/workdrive/test")
@permission_required("settings:write")
def workdrive_test():
    try:
        client = WorkDriveClient()
        folder = client.verify_root_folder()
        db().settings.update_one({"_id": "workdrive_oauth"}, {"$set": {"verified_at": now(), "verified_folder_id": folder["id"], "root_folder_name": folder["name"], "team_id": folder["team_id"]}, "$unset": {"verification_error": "", "error_kind": ""}}, upsert=True)
        return jsonify(ok=True, verified=True, verified_at=now(), root_folder_name=folder["name"])
    except WorkDriveError as exc:
        db().settings.update_one({"_id": "workdrive_oauth"}, {"$set": {"verification_error": str(exc), "error_kind": exc.kind, "failed_at": now()}}, upsert=True)
        return jsonify(error=str(exc), status=exc.kind), 503


@api.post("/workdrive/disconnect")
@permission_required("settings:write")
def workdrive_disconnect():
    db().settings.delete_one({"_id": "workdrive_oauth"})
    return jsonify(ok=True)


def _mail_status_payload():
    client = ZohoMailClient()
    stored = client.record
    verified = bool(stored.get("verified_at") and stored.get("account_id") and client.encrypted_refresh_token)
    if verified:
        status = "connected"
    elif "token_encryption_key" in client.missing:
        status = "missing_encryption_key"
    elif "refresh_token" in client.missing:
        status = "missing_refresh_token"
    else:
        status = stored.get("error_kind") or "incomplete"
    settings = db().settings.find_one({"_id": "zoho_mail_settings"}) or {}
    senders = stored.get("sender_addresses", []) if verified else []
    if verified and not senders and stored.get("account_email"):
        senders = [{"address": stored["account_email"], "display_name": ""}]
    selected_sender = settings.get("sender_address") or (stored.get("account_email") if verified else None)
    return {"configured": client.configured, "missing": client.missing, "verified": verified, "status": status, "verification_error": stored.get("verification_error"), "account_email": stored.get("account_email") if verified else None, "account_id": stored.get("account_id") if verified else None, "sender_addresses": senders, "organization": stored.get("organization") if verified else None, "connected_at": stored.get("connected_at") if verified else None, "verified_at": stored.get("verified_at") if verified else None, "last_sent_at": stored.get("last_sent_at") if verified else None, "scopes": [scope.strip() for scope in client.scopes.split(",") if scope.strip()], "email_settings": {"sender_address": selected_sender or "", "sender_display_name": settings.get("sender_display_name", ""), "reply_to_email": settings.get("reply_to_email", ""), "default_signature": settings.get("default_signature", ""), "notifications_enabled": bool(settings.get("notifications_enabled", False))}}


@api.get("/mail/status")
@auth_required
def mail_status():
    return jsonify(_mail_status_payload())


@api.post("/mail/connect")
@permission_required("settings:write")
def mail_connect():
    try:
        state = str(uuid4())
        db().settings.insert_one({"_id": f"mail_state:{state}", "state": state, "user_id": g.user["_id"], "expires_at": now() + timedelta(minutes=10), "created_at": now()})
        return jsonify(authorization_url=ZohoMailClient().authorization_url(state))
    except WorkDriveError as exc:
        return jsonify(error=str(exc), status=exc.kind), 503


@api.post("/mail/disconnect")
@permission_required("settings:write")
def mail_disconnect():
    db().settings.delete_one({"_id": ZohoMailClient.record_id})
    return jsonify(ok=True)


@api.post("/mail/test")
@permission_required("settings:write")
def mail_test():
    try:
        account = ZohoMailClient().verify_account()
        return jsonify(ok=True, verified=True, account_email=account.get("email"), organization=account.get("organization"), verified_at=now())
    except WorkDriveError as exc:
        db().settings.update_one({"_id": ZohoMailClient.record_id}, {"$set": {"verification_error": str(exc), "error_kind": exc.kind, "failed_at": now()}}, upsert=True)
        return jsonify(error=str(exc), status=exc.kind), 503


@api.patch("/mail/settings")
@permission_required("settings:write")
def mail_settings_update():
    data = body()
    allowed = {"sender_address", "sender_display_name", "reply_to_email", "default_signature", "notifications_enabled"}
    changes = {key: data[key] for key in allowed if key in data}
    if "reply_to_email" in changes and changes["reply_to_email"] and "@" not in changes["reply_to_email"]:
        return jsonify(error="Reply-to email address is invalid"), 400
    if "sender_address" in changes:
        oauth = ZohoMailClient().record
        allowed_senders = {item.get("address", "").lower() for item in oauth.get("sender_addresses", [])}
        if oauth.get("account_email"):
            allowed_senders.add(oauth["account_email"].lower())
        if not changes["sender_address"] or changes["sender_address"].lower() not in allowed_senders:
            return jsonify(error="Select a sender address verified for the connected Zoho Mail account"), 400
    changes["updated_at"] = now()
    changes["updated_by"] = g.user["_id"]
    db().settings.update_one({"_id": "zoho_mail_settings"}, {"$set": changes}, upsert=True)
    return jsonify(ok=True, settings=_mail_status_payload()["email_settings"])


@api.post("/mail/test-email")
@permission_required("settings:write")
def mail_test_email():
    try:
        data = body(("to",))
        if "@" not in data["to"]:
            return jsonify(error="Recipient email address is invalid"), 400
        settings = db().settings.find_one({"_id": "zoho_mail_settings"}) or {}
        result = ZohoMailClient().send_message(data["to"], "RK Fashion portal test email", "<p>Your Zoho Mail connection to the RK Fashion Operations Portal is working.</p>", settings.get("sender_display_name"), settings.get("reply_to_email"))
        return jsonify(ok=True, accepted=result["accepted"])
    except (ValueError, WorkDriveError) as exc:
        kind = exc.kind if isinstance(exc, WorkDriveError) else "validation_error"
        return jsonify(error=str(exc), status=kind), 503 if isinstance(exc, WorkDriveError) else 400


@api.get("/integrations/zoho/mail/callback")
def mail_callback():
    frontend = current_app.config["FRONTEND_ORIGIN"]
    state = request.args.get("state", "")
    state_doc = db().settings.find_one({"_id": f"mail_state:{state}", "state": state})
    expires_at = utc_datetime(state_doc.get("expires_at")) if state_doc else None
    if not state_doc or (expires_at or utc_datetime(now())) < utc_datetime(now()):
        return redirect(f"{frontend}/workspace/mail?mail=error&message=Invalid%20or%20expired%20OAuth%20state")
    db().settings.delete_one({"_id": state_doc["_id"]})
    if request.args.get("error"):
        return redirect(f"{frontend}/workspace/mail?mail=error&message=Zoho%20authorization%20was%20cancelled")
    try:
        client = ZohoMailClient()
        client.exchange_code(request.args.get("code", ""))
        client.verify_account()
        db().settings.update_one({"_id": client.record_id}, {"$unset": {"verification_error": "", "error_kind": ""}, "$set": {"verified_by": state_doc["user_id"]}}, upsert=True)
        return redirect(f"{frontend}/workspace/mail?mail=connected")
    except WorkDriveError as exc:
        db().settings.update_one({"_id": ZohoMailClient.record_id}, {"$set": {"verification_error": str(exc), "error_kind": exc.kind, "failed_at": now()}, "$unset": {"verified_at": "", "account_id": ""}}, upsert=True)
        return redirect(f"{frontend}/workspace/mail?mail=error&message={quote(str(exc))}")


@api.post("/workdrive/connect")
@permission_required("settings:write")
def workdrive_connect():
    try:
        state = str(uuid4())
        db().settings.insert_one({"_id": f"workdrive_state:{state}", "state": state, "user_id": g.user["_id"], "expires_at": now() + timedelta(minutes=10), "created_at": now()})
        return jsonify(authorization_url=WorkDriveClient().authorization_url(state))
    except WorkDriveError as exc:
        return jsonify(error=str(exc)), 503


@api.get("/integrations/zoho/callback")
def workdrive_callback():
    frontend = current_app.config["FRONTEND_ORIGIN"]
    state = request.args.get("state", "")
    state_doc = db().settings.find_one({"_id": f"workdrive_state:{state}", "state": state})
    expires_at = utc_datetime(state_doc.get("expires_at")) if state_doc else None
    if not state_doc or (expires_at or utc_datetime(now())) < utc_datetime(now()):
        return redirect(f"{frontend}/workspace/settings?workdrive=error&message=Invalid%20or%20expired%20OAuth%20state")
    db().settings.delete_one({"_id": state_doc["_id"]})
    if request.args.get("error"):
        return redirect(f"{frontend}/workspace/settings?workdrive=error&message=Zoho%20authorization%20was%20cancelled")
    try:
        client = WorkDriveClient()
        client.exchange_code(request.args.get("code", ""))
        folder = client.discover_root_folder()
        connected_at = (db().settings.find_one({"_id": "workdrive_oauth"}) or {}).get("connected_at") or now()
        db().settings.update_one({"_id": "workdrive_oauth"}, {"$set": {"root_folder_id": folder["id"], "root_folder_name": folder["name"], "team_id": folder["team_id"], "connected_at": connected_at, "verified_at": now(), "verified_by": state_doc["user_id"], "verified_folder_id": folder["id"]}, "$unset": {"verification_error": "", "error_kind": ""}}, upsert=True)
        return redirect(f"{frontend}/workspace/settings?workdrive=connected")
    except WorkDriveError as exc:
        db().settings.update_one({"_id": "workdrive_oauth"}, {"$set": {"verification_error": str(exc), "error_kind": exc.kind, "failed_at": now()}, "$unset": {"verified_at": "", "verified_folder_id": ""}}, upsert=True)
        return redirect(f"{frontend}/workspace/settings?workdrive=error&message={quote(str(exc))}")


@api.post("/documents/<document_id>/sync-workdrive")
@permission_required("documents:write")
def sync_document_workdrive(document_id):
    document = db().documents.find_one({"_id": oid(document_id)})
    if not document:
        return jsonify(error="Document not found"), 404
    if document.get("workdrive_file_id") and document.get("upload_status") == "synced":
        return jsonify(document=serialize(document), duplicate_prevented=True)
    if not document.get("stored_path") or not os.path.exists(document["stored_path"]):
        return jsonify(error="This generated document must be downloaded and retained locally before synchronization"), 409
    db().documents.update_one({"_id": document["_id"]}, {"$set": {"upload_status": "upload_pending", "sync_attempted_at": now()}})
    try:
        result = WorkDriveClient().upload(document["stored_path"], document.get("workdrive_folder_id"), document.get("original_name"))
        file_data = (result.get("data") or [{}])[0]
        attributes = file_data.get("attributes", {})
        db().documents.update_one({"_id": document["_id"]}, {"$set": {"upload_status": "synced", "storage_provider": "local_and_workdrive", "workdrive_file_id": file_data.get("id"), "workdrive_link": attributes.get("permalink"), "synced_at": now()}, "$unset": {"sync_error": ""}})
    except WorkDriveError as exc:
        db().documents.update_one({"_id": document["_id"]}, {"$set": {"upload_status": "failed", "sync_error": str(exc), "updated_at": now()}})
        return jsonify(error=str(exc), status="failed"), 503
    return jsonify(document=serialize(db().documents.find_one({"_id": document["_id"]})))
